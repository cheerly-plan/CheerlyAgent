"""知识库 P1：权限守卫、检索工具注册与切片页码溯源。"""

import inspect
import tempfile
from datetime import UTC, datetime
from pathlib import Path

import pytest
from fastapi import HTTPException

from ai_agent.api.dependencies import get_current_user, require_admin
from ai_agent.api.routers.auth import _to_auth_response, me
from ai_agent.api.routers.kb import (
    _clean_tags,
    delete_file,
    download_file,
    edit_chunk,
    init_upload,
    list_chunks,
    list_files,
    rename_file,
    set_file_tags,
)
from ai_agent.api.routers.kg import build_graphs, entity_chunks, file_graph, global_graph, rebuild_graph, search_graph
from ai_agent.api.routers.user import user_info
from ai_agent.core.config import settings
from ai_agent.kb import vision
from ai_agent.kb.chunking import chunk_segments, chunk_text
from ai_agent.kb.parsing import Segment, parse_segments, supported
from ai_agent.kb.pipeline import KbPipeline
from ai_agent.memory.db import UserRecord
from ai_agent.tools.kb_search import _format_hit
from ai_agent.tools.registry import TOOLS, tool_label, tool_schemas


def _user(username: str) -> UserRecord:
    return UserRecord(user_id="uid-1", username=username, display_name=username)


def test_require_admin_allows_whitelisted(monkeypatch):
    monkeypatch.setattr(settings, "admin_usernames", "alice, bob")
    assert require_admin(_user("bob")).username == "bob"


def test_require_admin_rejects_others(monkeypatch):
    monkeypatch.setattr(settings, "admin_usernames", "alice")
    with pytest.raises(HTTPException) as exc:
        require_admin(_user("mike"))
    assert exc.value.status_code == 403


def test_require_admin_empty_whitelist_denies_everyone(monkeypatch):
    # 空配置 = 没有管理员，写操作全 403（只读浏览不受影响）
    monkeypatch.setattr(settings, "admin_usernames", "")
    with pytest.raises(HTTPException):
        require_admin(_user("alice"))


def test_user_info_reports_admin_flag(monkeypatch):
    monkeypatch.setattr(settings, "admin_usernames", "alice")
    assert user_info(_user("alice")).is_admin is True
    assert user_info(_user("mike")).is_admin is False


def test_login_response_reports_admin_flag(monkeypatch):
    # 登录响应直接带 is_admin，前端才能立即决定是否渲染知识库入口
    monkeypatch.setattr(settings, "admin_usernames", "alice")
    monkeypatch.setattr(
        "ai_agent.api.routers.auth.store.issue_token",
        lambda _uid: ("tok", datetime.now(UTC)),
    )
    assert _to_auth_response(_user("alice")).user.is_admin is True
    assert _to_auth_response(_user("mike")).user.is_admin is False


def test_me_reports_admin_flag(monkeypatch):
    # 登录响应、/v1/auth/me、/user/info 三个入口必须同构带上 is_admin，
    # 前端换任何一个来源取用户信息都不能丢知识库入口
    monkeypatch.setattr(settings, "admin_usernames", "alice")
    assert me(_user("alice")).is_admin is True
    assert me(_user("mike")).is_admin is False


def test_kb_reads_open_to_all_users():
    """读接口（列表/切片/下载/看图/图谱检索）对所有登录用户开放，普通用户也能查。"""
    for func, param in [
        (list_files, "_user"),
        (list_chunks, "_user"),
        (download_file, "_user"),
        (global_graph, "_user"),
        (file_graph, "_user"),
        (search_graph, "_user"),
        (entity_chunks, "_user"),
    ]:
        dep = inspect.signature(func).parameters[param].default
        assert dep.dependency is get_current_user, func.__name__


def test_kb_writes_require_admin():
    """写操作（上传建档/删除/改标签/切片修订/建图/重建）仍然卡在管理员身后。"""
    for func, param in [
        (init_upload, "user"),
        (delete_file, "_user"),
        (set_file_tags, "_user"),
        (edit_chunk, "user"),
        (build_graphs, "user"),
        (rebuild_graph, "user"),
    ]:
        dep = inspect.signature(func).parameters[param].default
        assert dep.dependency is require_admin, func.__name__


def test_kb_chunk_edit_requires_admin():
    # 切片修订会改写检索库内容，必须卡在管理员身后
    dep = inspect.signature(edit_chunk).parameters["user"].default
    assert dep.dependency is require_admin


def test_kg_build_requires_admin():
    # 批量建图触发 LLM 抽取，同样只对管理员开放
    dep = inspect.signature(build_graphs).parameters["user"].default
    assert dep.dependency is require_admin


def test_pipeline_does_not_autobuild_graph():
    # 上传后不再自动建图：图谱由管理员核对切片后选文件显式触发，
    # 防止「建完又改切片 → 整份重抽」的浪费
    source = inspect.getsource(KbPipeline._process)
    assert "build_file_graph" not in source
    assert "_build_graph" not in source


def test_kb_search_registered():
    names = [item["function"]["name"] for item in tool_schemas()]
    assert "kb_search" in names
    assert "kb_search" in [tool.name for tool in TOOLS]
    assert tool_label("kb_search") == "知识库检索"


def test_format_hit_uses_page_for_pdf():
    assert _format_hit(
        {"filename": "手册.pdf", "page": 3, "seq": 7, "text": "正文"}
    ) == "[手册.pdf · 第3页]\n正文"


def test_format_hit_falls_back_to_seq_without_page():
    assert _format_hit(
        {"filename": "笔记.md", "page": None, "seq": 0, "text": "正文"}
    ) == "[笔记.md · 第1段]\n正文"


def test_chunk_segments_keeps_page_of_each_chunk():
    segments = [
        Segment("第一页内容。" * 30, 1),
        Segment("第二页内容。", 2),
        Segment("无页文档"),
    ]
    chunks = chunk_segments(segments, max_chars=100, overlap=10)
    # 第 1 页长文拆成多片，页码全部保留为 1；页面不跨页合并
    assert [c.page for c in chunks][:2] == [1, 1]
    assert chunks[-2].page == 2
    assert chunks[-1].page is None
    assert all(c.text for c in chunks)


def test_chunk_text_short_input_untouched():
    assert chunk_text("你好世界", max_chars=500, overlap=50) == ["你好世界"]


def test_chunk_text_blank_returns_empty():
    assert chunk_text("   \n\n ") == []


# ——— 表格与图表：内容类型化 ———


def test_docx_table_is_extracted_as_markdown():
    # 之前的实现只遍历 paragraphs，Word 里的数据表整表丢失
    from docx import Document

    doc = Document()
    doc.add_paragraph("2024 年营收同比增长 18%。")
    table = doc.add_table(rows=3, cols=3)
    rows = [
        ["年份", "营收(亿)", "增速"],
        ["2023", "120", "12%"],
        ["2024", "142", "18%"],
    ]
    for r, row in enumerate(rows):
        for c, value in enumerate(row):
            table.cell(r, c).text = value
    path = Path(tempfile.mkdtemp()) / "年报.docx"
    doc.save(path)

    segments = parse_segments(path)
    assert any(seg.kind == "table" for seg in segments)
    table_text = next(seg.text for seg in segments if seg.kind == "table")
    assert "| 年份 | 营收(亿) | 增速 |" in table_text
    assert "| 2024 | 142 | 18% |" in table_text


def test_xlsx_sheet_becomes_table_segment():
    from openpyxl import Workbook

    workbook = Workbook()
    sheet = workbook.active
    sheet.title = "营收"
    sheet.append(["年份", "营收(亿)"])
    sheet.append([2023, 120])
    sheet.append([2024, 142.0])
    path = Path(tempfile.mkdtemp()) / "数据.xlsx"
    workbook.save(path)

    segments = parse_segments(path)
    assert len(segments) == 1
    seg = segments[0]
    assert seg.kind == "table"
    assert seg.page == 1  # sheet 序号，展示为「表1」
    assert "工作表：营收" in seg.text
    assert "| 2024 | 142 |" in seg.text  # 142.0 归一成 142


def test_xlsx_supported_suffix():
    assert supported(".xlsx")
    assert supported(".XLSM")
    assert not supported(".xls")


def test_chunk_table_kept_whole_when_short():
    text = "| 年份 | 营收 |\n|---|---|\n| 2024 | 142 |"
    chunks = chunk_segments([Segment(text, 1, "table")], max_chars=500)
    assert len(chunks) == 1
    assert chunks[0].kind == "table"
    assert chunks[0].page == 1


def test_chunk_long_table_repeats_header():
    rows = "\n".join(f"| {year} | {year * 10} |" for year in range(2000, 2060))
    text = "| 年份 | 营收 |\n|---|---|\n" + rows
    chunks = chunk_segments([Segment(text, 1, "table")], max_chars=200)
    assert len(chunks) > 1
    # 每片都带表头，否则行列对应关系丢失
    assert all(c.text.startswith("| 年份 | 营收 |") for c in chunks)
    assert all(c.kind == "table" for c in chunks)


def test_format_hit_labels_content_kind():
    base = {"filename": "报告.pdf", "seq": 0, "text": "内容"}
    assert _format_hit({**base, "kind": "text", "page": 3}) == "[报告.pdf · 第3页]\n内容"
    assert _format_hit({**base, "kind": "table", "page": 2}) == "[报告.pdf · 表2]\n内容"
    figure = _format_hit({**base, "kind": "figure", "page": 5})
    assert "第5页图表" in figure
    assert "图像识别" in figure


def test_vision_disabled_by_default(monkeypatch):
    # 未配置模型即关闭：默认不产生任何视觉调用费用
    monkeypatch.setattr(settings, "kb_vision_model", "")
    assert vision.enabled() is False
    assert vision.describe_image(b"fake-png") is None


# ——— 管理能力：标签 / 内容去重 / 下载 ———


def test_clean_tags_strips_dedupes_and_caps():
    assert _clean_tags(["  考勤 ", "", "考勤", "制" * 30]) == ["考勤", "制" * 24]
    assert _clean_tags([f"t{i}" for i in range(20)]) == [f"t{i}" for i in range(8)]
    assert _clean_tags([]) == []


def _prepare_complete(request, monkeypatch):
    """complete_upload 的公共桩：建档记录 + 分片落盘 + store/pipeline 打桩。

    临时目录放项目内（.tmp-kb-test）：系统 tmp 目录在沙箱里写不进去。
    """
    import shutil
    import uuid as _uuid

    from ai_agent.api.routers import kb as kb_router

    base = Path(__file__).resolve().parent.parent / ".tmp-kb-test"
    if not base.exists():
        base.mkdir()
    upload_root = base / _uuid.uuid4().hex
    record = {
        "id": _uuid.uuid4(),
        "filename": "a.pdf",
        "size_bytes": 4,
        "status": "uploaded",
    }
    monkeypatch.setattr(kb_router, "_get_file", lambda fid: record)
    monkeypatch.setattr(kb_router.settings, "kb_upload_dir", upload_root)
    part_dir = upload_root / str(record["id"]) / "parts"
    part_dir.mkdir(parents=True)
    (part_dir / "000000.part").write_bytes(b"ABCD")
    # 测试收尾清理整个临时根目录（record 目录已由被测代码自己删）
    request.addfinalizer(lambda: shutil.rmtree(base, ignore_errors=True))
    return kb_router, record


def test_check_duplicate_rejects_same_content(monkeypatch):
    """同内容文件拒绝入库：409 + 删档案，重复检查在解析管线之前。"""
    import uuid as _uuid

    from ai_agent.api.routers import kb as kb_router

    fid = _uuid.uuid4()
    monkeypatch.setattr(
        kb_router.store,
        "find_kb_file_by_hash",
        lambda h, *, exclude_file_id=None: {
            "file_id": _uuid.uuid4(),
            "filename": "已有.pdf",
        },
    )
    removed: list = []
    monkeypatch.setattr(kb_router.store, "delete_kb_file", removed.append)
    monkeypatch.setattr(kb_router.shutil, "rmtree", lambda *a, **k: None)

    with pytest.raises(HTTPException) as exc:
        kb_router._check_duplicate("a" * 64, fid, "a.pdf", "alice")

    assert exc.value.status_code == 409
    assert "已有.pdf" in exc.value.detail
    assert removed == [fid]


def test_check_duplicate_passes_new_content(monkeypatch):
    import uuid as _uuid

    from ai_agent.api.routers import kb as kb_router

    monkeypatch.setattr(
        kb_router.store, "find_kb_file_by_hash", lambda h, *, exclude_file_id=None: None
    )
    deleted: list = []
    monkeypatch.setattr(kb_router.store, "delete_kb_file", deleted.append)

    # 不抛异常即通过
    kb_router._check_duplicate("a" * 64, _uuid.uuid4(), "a.pdf", "alice")
    assert deleted == []



# ---------- 改名 ----------


def _rename_record(request, monkeypatch, status="ready", filename="旧报告.pdf"):
    """rename_file 的公共桩：ready 记录 + 磁盘上放好旧名文件 + pipeline 打桩。"""
    import shutil
    import uuid as _uuid

    from ai_agent.api.routers import kb as kb_router

    base = Path(__file__).resolve().parent.parent / ".tmp-kb-test"
    if not base.exists():
        base.mkdir()
    upload_root = base / _uuid.uuid4().hex
    fid = _uuid.uuid4()
    record = {"id": fid, "filename": filename, "size_bytes": 4, "status": status}
    monkeypatch.setattr(kb_router, "_parse_file_id", lambda _fid: fid)
    monkeypatch.setattr(kb_router, "_get_file", lambda _fid: record)
    monkeypatch.setattr(kb_router.settings, "kb_upload_dir", upload_root)
    renamed: list = []
    monkeypatch.setattr(
        kb_router.store,
        "rename_kb_file",
        lambda _fid, name: renamed.append(name) or True,
    )
    submitted: list = []
    monkeypatch.setattr(
        kb_router.kb_pipeline, "submit_reembed", lambda _fid: submitted.append(_fid)
    )
    request.addfinalizer(lambda: shutil.rmtree(base, ignore_errors=True))
    return kb_router, record, renamed, submitted


def test_rename_rejects_suffix_change(request, monkeypatch):
    from ai_agent.api.schemas import KbRenameRequest

    kb_router, record, renamed, submitted = _rename_record(request, monkeypatch)
    with pytest.raises(HTTPException) as exc:
        rename_file("x", KbRenameRequest(filename="新报告.docx"), _admin())
    assert exc.value.status_code == 400
    assert renamed == [] and submitted == []


def test_rename_requires_ready(request, monkeypatch):
    from ai_agent.api.schemas import KbRenameRequest

    kb_router, record, renamed, submitted = _rename_record(
        request, monkeypatch, status="parsing"
    )
    with pytest.raises(HTTPException) as exc:
        rename_file("x", KbRenameRequest(filename="新报告.pdf"), _admin())
    assert exc.value.status_code == 409
    assert renamed == [] and submitted == []


def test_rename_updates_db_disk_and_reembeds(request, monkeypatch):
    """stem 变了：DB 改名 + 磁盘文件重命名 + 后台重嵌，三件事都发生。"""
    import uuid as _uuid

    from ai_agent.api.schemas import KbRenameRequest

    kb_router, record, renamed, submitted = _rename_record(request, monkeypatch)
    upload_dir = kb_router.Path(kb_router.settings.kb_upload_dir) / str(record["id"])
    upload_dir.mkdir(parents=True)
    (upload_dir / record["filename"]).write_bytes(b"ABCD")

    result = rename_file(
        str(record["id"]), KbRenameRequest(filename="2026 新报告.pdf"), _admin()
    )

    assert result == {"ok": True, "filename": "2026 新报告.pdf", "reembedding": True}
    assert renamed == ["2026 新报告.pdf"]
    assert (upload_dir / "2026 新报告.pdf").is_file()
    assert not (upload_dir / record["filename"]).exists()
    assert submitted == [record["id"]]


def test_rename_same_stem_skips_reembed(request, monkeypatch):
    """仅扩展名大小写变化（stem 不变）：向量语境没变，不值得整份重嵌。"""
    from ai_agent.api.schemas import KbRenameRequest

    kb_router, record, renamed, submitted = _rename_record(
        request, monkeypatch, filename="报告.PDF"
    )
    result = rename_file("x", KbRenameRequest(filename="报告.pdf"), _admin())
    assert result["reembedding"] is False
    assert renamed == ["报告.pdf"] and submitted == []


def _admin():
    return UserRecord(
        user_id="u1", username="alice", display_name="Alice",
        created_at=datetime.now(UTC).isoformat(),
    )


# ---------- 并行向量化与并行视觉 ----------


def test_embed_texts_parallel_keeps_order(monkeypatch):
    """多批并发执行，但结果顺序与输入严格一致（写错位会毁掉全部检索）。"""
    from ai_agent.memory import embeddings

    def fake_embed(texts):
        return [[len(t), 0] for t in texts]

    monkeypatch.setattr(embeddings, "embed_texts", fake_embed)
    texts = [f"t{i:02d}" for i in range(25)]  # 10/批 → 3 批
    out = embeddings.embed_texts_parallel(texts, concurrency=3)
    assert out == fake_embed(texts)


def test_embed_texts_parallel_single_batch_passthrough(monkeypatch):
    """不足两批直接走串行路径，不开线程池。"""
    from ai_agent.memory import embeddings

    called = []
    monkeypatch.setattr(
        embeddings, "embed_texts", lambda ts: called.append(list(ts)) or [[1]]
    )
    assert embeddings.embed_texts_parallel(["a"], concurrency=4) == [[1]]
    assert embeddings.embed_texts_parallel(["a", "b"], concurrency=0) == [[1]]
    assert embeddings.embed_texts_parallel([], concurrency=4) == []


def test_embed_texts_parallel_failure_propagates(monkeypatch):
    """任一批失败抛 LLMError：与串行版失败语义一致，调用方决定怎么兜。"""
    import pytest as _pytest

    from ai_agent.core.exceptions import LLMError
    from ai_agent.memory import embeddings

    def flaky(texts):
        if texts and texts[0] == "t10":
            raise LLMError("boom")
        return [[0] for _ in texts]

    monkeypatch.setattr(embeddings, "embed_texts", flaky)
    with _pytest.raises(LLMError):
        embeddings.embed_texts_parallel([f"t{i}" for i in range(20)], concurrency=3)


def test_describe_parallel_keeps_order(monkeypatch):
    """视觉并行：结果顺序与输入一致（页码归属不能乱）。"""
    from ai_agent.kb import parsing

    def fake_describe(blob, mime):
        return f"desc-{blob.decode()}"

    monkeypatch.setattr(parsing.vision, "describe_image", fake_describe)
    jobs = [(b"1", "image/png"), (b"2", "image/png"), (b"3", "image/png")]
    assert parsing._describe_parallel(jobs) == ["desc-1", "desc-2", "desc-3"]

    monkeypatch.setattr(parsing.settings, "kb_vision_concurrency", 1)
    assert parsing._describe_parallel(jobs) == ["desc-1", "desc-2", "desc-3"]


# ——— 图表页筛选：别把表格页当图表送视觉模型 ———


class _FakePage:
    """只实现 _page_has_chart 用到的那几个方法。"""

    def __init__(self, images=(), drawings=0, size=(600, 800)):
        self._images = list(images)
        self._drawings = drawings
        self.rect = type("Rect", (), {"width": size[0], "height": size[1]})()

    def get_image_info(self):
        return [{"bbox": (0, 0, w, h)} for w, h in self._images]

    def get_images(self, full=False):
        return [(1, 0, w, h, 8, "", "", "", "") for w, h in self._images]

    def get_drawings(self):
        if self._drawings == "boom":
            raise ValueError("坏页")
        return [{}] * self._drawings


def test_chart_page_skips_dense_text_page():
    """整页表格：几百条矢量边框 + 上千字正文 —— 不是图表页。

    实测 331 页招股书按「路径数 ≥ 150」误判 124 页，20 页预算全烧在
    目录和财务报表上，每页还要等一次视觉调用。
    """
    from ai_agent.kb import parsing

    page = _FakePage(drawings=215, images=[(80, 60)])
    assert parsing._page_has_chart(page, text_len=1085) is False


def test_chart_page_accepts_large_image_with_little_text():
    from ai_agent.kb import parsing

    page = _FakePage(images=[(500, 600)])  # 占页面大半
    assert parsing._page_has_chart(page, text_len=180) is True


def test_chart_page_ignores_logo_sized_image():
    """页眉 logo / 装饰线不算图表：面积占比够不上阈值。"""
    from ai_agent.kb import parsing

    page = _FakePage(images=[(60, 40)], drawings=3)
    assert parsing._page_has_chart(page, text_len=200) is False


def test_chart_page_accepts_vector_only_chart():
    """ECharts 走 SVG 导出时没有位图，靠路径数识别（但要求文字少）。"""
    from ai_agent.kb import parsing

    page = _FakePage(drawings=200)
    assert parsing._page_has_chart(page, text_len=120) is True


def test_chart_page_survives_broken_page():
    """单页取图信息报错不能把整份解析带崩。"""
    from ai_agent.kb import parsing

    page = _FakePage(drawings="boom")
    assert parsing._page_has_chart(page, text_len=100) is False


def test_rebuild_graph_reads_chunks_from_db(monkeypatch):
    """图谱重建不再重新解析源文件：PDF 重新解析要再渲染 + 视觉识别一遍，
    几分钟白花，而且重切的 seq 与库里的切片对不上。"""
    import uuid as _uuid

    from ai_agent.kb import pipeline as pl

    seen: dict = {}
    monkeypatch.setattr(
        pl.store,
        "list_kb_chunk_inputs",
        lambda fid: [
            {"id": 1, "text": "第一段正文，公司主营装饰工程。", "page": 1, "kind": "text"},
            {"id": 2, "text": "第二段正文，报告期内营收增长。", "page": 2, "kind": "text"},
        ],
    )
    monkeypatch.setattr(pl.store, "update_graph_status", lambda *a, **k: seen.setdefault("status", a[1]))
    monkeypatch.setattr(pl, "build_file_graph", lambda fid, chunks, **k: seen.update(chunks=chunks) or {"entities": 3})

    def boom(*a, **k):
        raise AssertionError("重建图谱不该重新解析源文件")

    monkeypatch.setattr(pl, "parse_segments", boom)

    pl.kb_pipeline._rebuild_graph(_uuid.uuid4())

    assert [c.text for c in seen["chunks"]] == [
        "第一段正文，公司主营装饰工程。",
        "第二段正文，报告期内营收增长。",
    ]
    assert [c.page for c in seen["chunks"]] == [1, 2]
