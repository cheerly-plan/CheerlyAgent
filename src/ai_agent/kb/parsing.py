"""文件解析：把上传的原始文件转成带类型的文本段。

支持 PDF / DOCX / XLSX / MD / TXT / CSV。产出 Segment 列表，每段带
kind（正文 / 表格 / 图表）与页码，切片时按 kind 分别处理：

- text：普通正文，PDF 按页拆分，循页码可溯源
- table：Word 表格、Excel 工作表。转成 Markdown 表格保留行列对应关系，
  切片时整表成一片、超长则重复表头（见 chunking）
- figure：PDF 含图页 / DOCX 图片经视觉模型转出的描述。ECharts 这类
  图表的数值不在文件里，只能靠读图近似还原，故单独标记，检索时提示
  「图像识别」以免被当成原文数据

解析失败统一抛 ParseError（带可读原因），由管道记入 file.error。
"""

import re
from concurrent.futures import ThreadPoolExecutor
from dataclasses import dataclass, replace
from io import BytesIO
from itertools import groupby
from pathlib import Path

from ai_agent.core.config import settings
from ai_agent.core.logging import get_logger
from ai_agent.kb import sections, vision

logger = get_logger(__name__)

TEXT = "text"
TABLE = "table"
FIGURE = "figure"


@dataclass(frozen=True)
class Segment:
    text: str
    # 页码语义：PDF 是页码，Excel 是 sheet 序号，其余为 None
    page: int | None = None
    kind: str = TEXT
    # 所属一级节（如「第十节 管理层讨论与分析」）；无法识别章节的文档为 None
    section: str | None = None
    # 报表口径（consolidated / parent），只对表格段有意义，见 _table_scope
    scope: str | None = None
    # 跨页/拆开后合并的表：起始页在 page，结束页在这里；单页表为 None
    page_end: int | None = None

_SUPPORTED = {".pdf", ".docx", ".xlsx", ".xlsm", ".md", ".markdown", ".txt", ".csv"}


class ParseError(Exception):
    """文件解析失败，消息面向用户展示。"""


def supported(suffix: str) -> bool:
    return suffix.lower() in _SUPPORTED


def supported_label() -> str:
    return "PDF / DOCX / XLSX / MD / TXT / CSV"


def parse_segments(path: Path) -> list[Segment]:
    suffix = path.suffix.lower()
    if suffix == ".pdf":
        return _parse_pdf(path)
    if suffix == ".docx":
        return _parse_docx(path)
    if suffix in {".xlsx", ".xlsm"}:
        return _parse_xlsx(path)
    if suffix in _SUPPORTED:
        return [
            Segment(path.read_text(encoding="utf-8", errors="replace"))
        ]
    raise ParseError(f"暂不支持的文件类型 {suffix}")


# ——— PDF ———


def _parse_pdf(path: Path) -> list[Segment]:
    segments, text_lens = _pdf_text(path)
    if not segments:
        # pymupdf 不可用 / 无文本层：退回整页抽取（无章节、无页眉剔除），
        # 扫描件走这条路时仍是「抽不出字就报错」，不做假解析
        segments, text_lens = _pdf_text_fallback(path)

    # 图表单独走渲染 + 视觉识别（pypdf 不能把页转成图）
    segments.extend(_pdf_figures(path, text_lens))

    if not segments:
        raise ParseError("PDF 中没有可提取的文本（可能是扫描件/图片版）")
    return segments


def _pdf_text(path: Path) -> tuple[list[Segment], list[int]]:
    """行级抽取：剔页眉页脚，按一级节分组后逐页产出 Segment；表格单独成段。

    返回 ([], []) 表示 pymupdf 不可用或文件打不开，由调用方降级——不用
    异常是因为「降级」是正常路径（未装 pymupdf 的部署仍要能解析）。

    这里只做「节 → 页」的拆分，节内跨页聚合留给 chunking：切片需要页码
    溯源，Segment 保持页粒度就不会丢页码。

    表格与正文混排时按 (页, 纵向位置) 排回原顺序——表格的位置决定它
    在哪一节、挨着哪段正文，统统挪到文末会让引用张冠李戴。
    """
    try:
        import pymupdf
    except ImportError:
        logger.warning("未安装 pymupdf，PDF 文本退回整页抽取（无章节识别）")
        return [], []
    try:
        doc = pymupdf.open(str(path))
    except Exception:
        logger.warning("pymupdf 打开 PDF 失败，退回整页抽取", exc_info=True)
        return [], []
    try:
        lines = sections.extract_lines(doc)
    except Exception:
        logger.warning("pymupdf 抽取文本行失败，退回整页抽取", exc_info=True)
        return [], []
    # 勾选框符号归一化要在任何按文本的比较（页眉去重、标题识别、表格剔除）
    # 之前做，否则同一符号在两处长得不一样，去重去不干净
    lines = normalize_choice_lines(lines)
    # 表格识别只是「多产出几段」，失败不该拖累正文解析
    tables: list[tuple[int, float, tuple, list[list[str]]]] = []
    page_heights: dict[int, float] = {}
    if settings.kb_pdf_table_extraction:
        try:
            tables, page_heights = _pdf_tables(doc)
        except Exception:
            logger.warning("PDF 表格识别失败，本次只出正文切片", exc_info=True)
        try:
            # 无框释义表（"X 指 Y"）find_tables 按线条识别不到，单独补一轮；
            # 传已有表格框做排除，页上残留框线的部分不会被出两份
            tables += _pdf_glossary_tables(
                doc, [(page, bbox) for page, _y, bbox, _rows in tables]
            )
        except Exception:
            logger.warning("PDF 释义表识别失败，跳过", exc_info=True)
    doc.close()
    if not lines:
        return [], []

    # 图表页判定用的是原始文字量（含页眉、含表格），与历史行为保持一致
    text_lens = _page_char_counts(lines)
    # 表格里的数字字号常小于正文且字符量极大，算正文字号（按量加权取众数）
    # 会把它拉偏，导致真标题被判成页眉——所以先剔掉表格行再定字号
    rest = _drop_table_lines(lines, tables)
    body = sections.body_size(rest)
    kept, stats = sections.strip_noise(rest, body)
    if tables:
        stats["表格"] = len(tables)
    if stats:
        logger.info("PDF 噪声剔除与表格 %s", stats)

    if not settings.kb_section_chunking:
        groups = [sections.Section("", "", start_page=kept[0].page if kept else 1, lines=list(kept))]
    else:
        groups = sections.group_by_section(kept, body)
        kept = [ln for sec in groups for ln in sec.lines]
    # 表格不在正文行里，只能靠同页的标题位置推断它属于哪一小节：
    # 表题/表头之下、下一个标题之上的表，才算这一节的
    markers = _section_markers(groups)

    # (页, 纵向位置, 段)：文本段与表格段按原版面顺序归并
    items: list[tuple[int, float, Segment]] = []
    for y, seg in _page_segments(groups):
        items.append((seg.page or 0, y, seg))
    # 表题接在表格正文前面（一行）：切成多片时每片都会带上它，
    # 「母公司利润表 + 营业收入 + 数值」因此能落在同一个切片里
    for page, y0, _bbox, rows, caption, page_end in _merge_continued_tables(
        _attach_captions(tables, kept), page_heights
    ):
        markdown = _rows_to_markdown(rows)
        if not markdown:
            continue
        # 单位/审计状态跟表题一起写进正文：长表切开后每一片都还知道
        # 「数字是百万元、这张表未经审计」，否则后半段只剩格子
        prefix = _table_prefix(caption, _table_meta_above(kept, page, y0))
        text = f"{prefix}\n{markdown}" if prefix else markdown
        items.append(
            (
                page,
                y0,
                Segment(
                    text,
                    page,
                    TABLE,
                    _label_at(markers, page, y0),
                    scope=_table_scope(caption),
                    page_end=page_end if page_end != page else None,
                ),
            )
        )
    items.sort(key=lambda it: (it[0], it[1]))
    return [seg for _page, _y, seg in items], text_lens


def _section_markers(
    groups: list[sections.Section],
) -> list[tuple[int, float, str]]:
    """(页, 纵坐标, 章节标签)：每个小节首行的位置就是这一节的起点。

    行内顺序即版面顺序，所以标记天然升序，查表走二分即可。用「点」坐标
    而不是 y_ratio，是因为表格的 y0 就是 pymupdf 的页面坐标，两者要能直接比。
    """
    out: list[tuple[int, float, str]] = []
    for sec in groups:
        label = sec.label
        if not label or not sec.lines:
            continue
        first = sec.lines[0]
        y = first.bbox[1] if first.bbox else first.y_ratio * 1000
        out.append((first.page, y, label))
    return out


def _label_at(
    markers: list[tuple[int, float, str]], page: int, y: float
) -> str | None:
    """查某个位置所属的章节：取它上方最近的那个标题，跨页继续沿用。"""
    best: str | None = None
    for m_page, m_y, label in markers:
        if m_page > page or (m_page == page and m_y > y + 1):
            break
        best = label
    return best


def _page_segments(
    groups: list[sections.Section],
) -> list[tuple[float, Segment]]:
    """按 (节, 页) 聚成 Segment，同时带回该段首行的纵坐标（与表格排序混排用）。

    一格小节一页一段：同一页里「（一）主要会计数据」和「（二）非经常性
    损益」各成一段，切片因此不会跨小节拼接，每片都能标到具体小节。
    """
    out: list[tuple[float, Segment]] = []
    for sec in groups:
        label = sec.label or None
        for page, page_lines in groupby(sec.lines, key=lambda ln: ln.page):
            group = list(page_lines)
            # 版面折行接回段落：PDF 每行一个 line，直接按行拼接会把
            # 「占总营」「收94.84%」切成两段，切片边界和引用都断在词中间
            paras = _merge_wrapped(group)
            # 勾选行正好落在小节开头时没有前置问题可接（「（二）非经常性
            # 损益项目和金额」自己成了新小节的首行），把小节标题当问题补上，
            # 否则模型读到的是一句没有提问对象的「☑（已勾选）适用」
            if label and paras and _is_choice_line(paras[0]):
                paras[0] = _attach_choice(label.split(" · ")[-1], paras[0])
            text = "\n".join(paras).strip()
            if not text:
                continue
            first = group[0]
            y = first.bbox[1] if first.bbox else first.y_ratio * 1000
            out.append((y, Segment(text, page, TEXT, label)))
    return out


def _segments_by_page(
    lines: list[sections.Line], *, section: str | None = None
) -> list[Segment]:
    """连续的同类行按页聚成 Segment（行序即页序，可直接 groupby）。"""
    out: list[Segment] = []
    for page, page_lines in groupby(lines, key=lambda ln: ln.page):
        text = "\n".join(ln.text for ln in page_lines).strip()
        if text:
            out.append(Segment(text, page, TEXT, section))
    return out


# 勾选框（单选/多选 selections）。PDF 里的方框和对勾常取自 Wingdings 一类的
# 符号字体：抽出来是**私用区字符**（实测 \uf052 = 已勾、\uf0a3 = 空框），或者
# 是写法各异的符号（√ ✓ ■ □ ☐ ○）。后果是选择题全部失读——模型看到的
# 「\uf052适用□不适用」跟噪声没区别，□ 和私用区字符它分不出哪个代表已勾，
# 于是「单选题选了哪项」从来没被正确读出来（季报里这类问法有十几处：
# 是否需追溯调整、是否经过审计、是否存在xxx情况）。
# 做法：先统一成标准 Unicode，再把状态写成中文——只换成 ☑/☐ 还不够，这两个
# 符号在别的语料里也被混用作纯装饰，写死中文标签最稳。
_CHECKED_CHARS = "\uf052✓✔√■●"
_UNCHECKED_CHARS = "\uf0a3□☐○"
_MARK_CHECKED = "☑"
_MARK_UNCHECKED = "☐"
_HAS_MARK = re.compile(f"[{re.escape(_CHECKED_CHARS + _UNCHECKED_CHARS)}]")

# 勾选行：整行只有若干「符号+选项」，没有别的（「☑（已勾选）适用 ☐（未勾选）
# 不适用」「□是 ☑否」）。判据必须要求「以勾选符开头」——只看长度的话，
# 「公司是否需追溯调整…会计数据：☐（未勾选）是 ☑（已勾选）否」这种已经带上
# 提问的行有 38 个字，会被误判成纯选项行，再被补一个小节标题当提问主体。
_MAX_CHOICE_LEN = 60


def normalize_choices(text: str) -> str:
    """把勾选符号归一化并标注选中状态。

    「\\uf052适用□不适用」→「☑（已勾选）适用 ☐（未勾选）不适用」。
    没有勾选符号的文本原样返回（不加任何标记），绝大多数行的开销为零。
    """
    if not text or not _HAS_MARK.search(text):
        return text
    table = {ord(c): _MARK_CHECKED for c in _CHECKED_CHARS}
    table.update({ord(c): _MARK_UNCHECKED for c in _UNCHECKED_CHARS})
    out = text.translate(table)
    out = out.replace(_MARK_CHECKED, f"{_MARK_CHECKED}（已勾选）")
    out = out.replace(_MARK_UNCHECKED, f"{_MARK_UNCHECKED}（未勾选）")
    # 原文本里两个选项紧挨着（「□适用√不适用」），留个空格免得读成一串
    return re.sub(r"(?<=\S)([☑☐])", r" \1", out)
    return out


def normalize_choice_lines(
    lines: list[sections.Line],
) -> list[sections.Line]:
    """逐行做勾选归一化。Line 是 frozen dataclass，只能整体 replace。"""
    out: list[sections.Line] = []
    for ln in lines:
        new = normalize_choices(ln.text)
        out.append(replace(ln, text=new) if new != ln.text else ln)
    return out


def _is_choice_line(text: str) -> bool:
    """是否是「整行只有选项、没有提问对象」的勾选行。"""
    if not text.startswith((_MARK_CHECKED, _MARK_UNCHECKED)):
        return False
    marks = sum(1 for ch in text if ch in (_MARK_CHECKED, _MARK_UNCHECKED))
    return (
        bool(marks) and marks <= 6 and len(text) <= _MAX_CHOICE_LEN and not _SENT_END.search(text)
    )


def _attach_choice(head: str, choice: str) -> str:
    """把勾选结果接到它的问题后面。

    季报里这些行一律紧跟提问（「持股5%以上股东…出借股份情况」+「☑不适用」），
    合并后模型才知道「不适用」回答的是什么。几何坐标不能当判据：同一问法的
    行有的排到了右边界被当前的折行规则接上，有的没排满就成了两行，结果同一
    份文档里一半合并一半分离，反而更难开除统一处。
    """
    if not head:
        return choice
    if head.endswith(("：", ":", "（", "(", "、", "，")) or _SENT_END.search(head):
        return head + choice
    return f"{head}：{choice}"


# 版面折行：上一行排到右边界、下一行从左边距起，且上一行没有句末标点，
# 那就是同一个自然段被排版折开。不接回去的代价很直观——安克的
# 「占总营」「收94.84%」、「持」「续加大研发」会被切成两条段落，切片边界
# 落在词中间，检索命中后引用出来是半句话，图谱也从半句里抽实体。
_SENT_END = re.compile(r"[。！？；：!?;:…」』”\"）)》】]$")
# 换块信号：编号标题、表题、注释、单位行。上一行即使没有句末标点，
# 后面跟着这些也不该接（「……具体情况如下」+「（1）充电储能类产品」）
_BLOCK_START = re.compile(
    r"^(第[一二三四五六七八九十百]+[节章]|[一二三四五六七八九十]+、|"
    r"（[一二三四五六七八九十\d]+）|\(\d+\)|\d+[、.．]|[①-⑳]|"
    r"表\s*\d|附表|单位[:：]|注[:：]|资料来源)"
)
_EDGE_SLACK = 6.0  # 点：判「贴左右边距」的容差，中文排版两端对齐会差几个点


def _merge_wrapped(lines: list[sections.Line]) -> list[str]:
    """把版面折行的行接回自然段，返回段落文本列表。

    只认同页内的折行：跨页的那一行是另一个版面，接起来会把页眉残留、
    页脚页码掺进正文（它们的坐标早就被剔除了，但相邻行仍在）。
    """
    boxes = [ln.bbox for ln in lines if ln.bbox]
    plain = [ln.text.strip() for ln in lines if ln.text.strip()]
    if len(boxes) < 2:
        return plain
    left = min(b[0] for b in boxes)
    right = max(b[2] for b in boxes)
    slack = max((right - left) * 0.05, _EDGE_SLACK)

    paras: list[str] = []
    prev: sections.Line | None = None
    for ln in lines:
        text = ln.text.strip()
        if not text:
            continue
        if prev is not None and paras:
            if _is_choice_line(text):
                # 勾选行一律并入上一行，不靠几何判据（见 _attach_choice）
                paras[-1] = _attach_choice(paras[-1], text)
            elif _wraps(prev, ln, left, right, slack):
                paras[-1] += text
            else:
                paras.append(text)
        else:
            paras.append(text)
        prev = ln
    return paras


def _wraps(
    prev: sections.Line,
    cur: sections.Line,
    left: float,
    right: float,
    slack: float,
) -> bool:
    if not prev.bbox or not cur.bbox:
        return False
    if prev.bbox[2] < right - slack:
        return False  # 上一行没排满 → 本来就是段末
    if abs(cur.bbox[0] - left) > slack:
        return False  # 下一行缩进/另有栏位 → 不是同一段
    # 上一行也必须从左边距起（允许 2 字缩进）：封面上的「证券代码：…」
    # 与右侧「公告编号：…」排在同一行，前者贴左、后者贴右，不看 x0
    # 会把上下两栏的文字接成一句
    if prev.bbox[0] > left + slack * 4:
        return False
    if _SENT_END.search(prev.text.strip()):
        return False
    if len(prev.norm) <= 40 and _BLOCK_START.match(prev.norm):
        return False  # 上一行本身是标题，别把标题并进正文
    return not _BLOCK_START.match(cur.norm)


# ——— PDF 表格 ———
# PDF 里没有「表格」这个对象，只有画出来的线和摆好的字。pymupdf 的
# find_tables 靠线条（无框表格靠文字对齐）反推单元格，抽出来的是二维
# 字符串数组，转成 Markdown 表格交给切片层（整表成片、超长重复表头）。

_CELL_WS = re.compile(r"\s+")
# 单元格内的换行是排版断行：数字常被拆成 "5176 63\n."（小数点另起一行）
_NUM_DECIMAL = re.compile(r"(\d[\d,]*)\s+(\d+)\s*\.\s*(\d*)")
# 千分位也逃不过断行："114 376 427 06 , , ." 是 114,376,427.06
# （PDF 把逗号和小数点排到单元格末尾，前面只剩一串数字碎片）
_NUM_THOUSANDS = re.compile(r"^(\d+(?:\s+\d+)+)\s*(?:[,]\s*)*\.\s*(\d*)$")
# "2006年 1 6月 -" 是 2006年1-6月：连字符被甩到单元格最后
_NUM_RANGE = re.compile(r"(\d+)\s+(\d+月)\s*-\s*$")
# 中文字之间的空格也是排版字距（"木粉 复合油 ，"），不是真空格
_CJK = "\u4e00-\u9fff\u3040-\u30ff"
_CJK_SPACE = re.compile(f"([{_CJK}])\\s+([{_CJK}])")
_CJK_RE = re.compile(f"[{_CJK}]")
# 标点被单独排版时前后会多出空格："复合油 ，" / "（ 注"
_SPACE_BEFORE_PUNCT = re.compile(r"\s+([，。、；：？！）】》」』%])")
_SPACE_AFTER_PUNCT = re.compile(r"([（【《「『])\s+")
# 表头续行（"元）"、"（%"）都很短，超过这个长度基本是数据行不是续行
_MAX_HEADER_CELL = 6
# 纯数字（允许千分位、百分比、负号）：表头碎片不会是这个形状
_NUM_ONLY = re.compile(r"^[\d.,%+\-]+$")
# 跨页续表：表格顶边在这个纵坐标以内就算"从页顶开始"
_HEAD_CARRY_Y = 130.0
# 表题离表体多远以内还算它的标题（pt）。财务报表的「母公司利润表」离表体
# 约 58pt，中间还夹着「2026 年1—9 月」「编制单位：」「单位：」三行，
# 窗口太窄就只找到元信息行
_CAPTION_WINDOW = 120.0
# 表题候选里要跳过的是表格元信息，不是表名
_CAPTION_NOISE = re.compile(
    r"^(单位|币种|金额单位|编制单位|审计类型|数据来源|资料来源|注\s*[:：])"
    r"|^[√□■☑]|^(适用|不适用)|^\d{4}\s*年|^\d+\s*/\s*\d+$"
)
# 表题的结构：以「表」结尾（合并资产负债表/母公司利润表）、带（续）、
# 或以节编号开头（（一）主要会计数据和财务指标）。只认这三种，宁缺毋滥——
# 认错表题比没有表题更坏，会把别的表的主体安到这张表上
_CAPTION_SHAPE = re.compile(
    r"(表$|（续）$"
    r"|^\(?[一二三四五六七八九十]+[、.]\s*[^\d\s]"
    r"|^（[一二三四五六七八九十]+）"
    # 编号后面必须接非数字：「1、截至本报告期末」是标题，「3.1%，较上年…」是句子
    r"|^\d{1,2}\s*[、.．]\s*[^\d\s])"
)
# 以这些标点收尾的是句子/引导语，不是表题：「…如下：」「…提升0.2 个百分点。」
_CAPTION_SENTENCE = re.compile(r"[。：；，、]$")
# 表题的长度上限：再长就是正文句子了
_CAPTION_MAX_CHARS = 40
# 上一页最后一张表的底边离页底不足这么远，说明它是被页面切断的、不是自己结束的
_PAGE_BOTTOM_SLACK = 90.0
# 同页被识别器拆成两截时，两框之间允许的最大空隙。再大更像上下两张独立表
_SAME_PAGE_GAP = 72.0
# 列名单元格的字数上限：超过基本都是解释性长句
_HEADER_CELL_MAX = 20
# 带这些标点的是句子，不是列名。刻意不含「、」「，」——「质押、标记或冻结
# 情况」这种真列名就用顿号，排掉它反而漏认
_HEADER_SENTENCE = re.compile(r"[。；！？]")
# 另一种断法：整数连写、小数点/逗号被画成独立块甩到单元格末尾
# 「401825 , .」= 4,018.25、「5358 .」= 53.58、「人民币100元 .」= 人民币1.00元
# （有小数点时末两位是小数，只有逗号时是纯千分位）
_NUM_TAIL_BOTH = re.compile(r"^([^\d]{0,8})(\d{2,})([^\d]{0,8})\s*,\s*\.\s*$")
_NUM_TAIL_DOT = re.compile(r"^([^\d]{0,8})(\d{2,})([^\d]{0,8})\s*\.\s*$")
_NUM_TAIL_COMMA = re.compile(r"^([^\d]{0,8})(\d{4,})([^\d]{0,8})\s*,\s*$")
# 逗号被单独排版时连成一片："总股本 ,， 比例" 只留一个中文逗号
_MULTI_COMMA = re.compile(r"\s*[，,](\s*[，,])+")
# 逗号、负号被画成独立块甩在数字后面（还常一起出现）：
# "471,784,966.36 ," → 471,784,966.36；"30,360,622.17 - ," → -30,360,622.17
_NUM_TRAIL_SIGNS = re.compile(r"^([\d.,]*\d)\s*(-)?\s*,?\s*$")
# 日期的连字符被单独排版："2016 12 31 - -" 是 2016-12-31
_DATE_DASH = re.compile(r"(\d{4})\s+(\d{1,2})\s+(\d{1,2})\s*-\s*-")
# 商标有效期："20090807 . ." 是 2009.08.07（小数点被画成独立块）
_DATE_DOTS = re.compile(r"(\d{4})(\d{2})(\d{2})\s*\.\s*\.\s*")
# 表头日期月份不补零："2020630 . ." 是 2020.6.30（8 位的走上面，避免 20191231 被切成 2019.1.231）
_DATE_DOTS_MD = re.compile(r"(\d{4})(\d)(\d{2})\s*\.\s*\.\s*")
# 标点后的空格："总股本 ， 比例"
_SPACE_AFTER_PUNCT_2 = re.compile(r"([，。、；：？！%])\s+")
# 单位括号/百分号/数字后甩出来的孤立逗号：
# 「资产负债率（母公司%），」「1,031,675,913.97，」
_TRAIL_COMMA = re.compile(r"(?<=[）)%\d])\s*[，,]$")
# 千分位逗号被甩到后面，落在两个中文之间："换向器 , 项目"
# （只认英文逗号，中文逗号多半是真标点）
_CJK_COMMA = re.compile(f"(?<=[{_CJK}])\\s*,\\s*(?=[{_CJK}])")
# 合计行常把标签写在表头列上（和「项目」同一列居中），不能当成「表头位出现了真数据」
_TOTAL_LABEL = re.compile(r"合计|小计|总计")
# 整行只剩标点的「孤儿行」：千分位逗号、句读被 PDF 排成独立块插在多行
# 文本中间（", ，" / ". 。 ，"）。夹着英文点/逗号的是数字碎片，位置
# 不可考，删掉；纯中文标点（「， 、 ；」）能插回上一行的空格缝。
_PUNCT_ROW = re.compile(
    r"^[\s.,;:!？?%\-—–…()（）【】\[\]《》「」、，。；：\"'“”‘’]+$"
)
# 职责表这类长单元格：正文一行、标点单独一行。只含中文标点才敢回填，
# 英文逗号/点可能是千分位或小数点（见 _DOT_ROW / 发行概况表）。
_WRAP_PUNCT = re.compile(r"^[\s]*[，。、；：！？][\s，。、；：！？]*$")
_WRAP_PUNCT_MARK = re.compile(r"[，。、；：！？]")
# 只有一个点的孤儿行（前面可带千分位逗号）：是数字被甩出的小数点
_DOT_ROW = re.compile(r"^,?\s*\.$")
# 被甩出小数点的数字串：多段连写（"3 672 71"、"114 376 427 06"——末段是
# 小数，前面是整数部分），或单个整数（"365"、"200000"）。
# 前后不能挨着数字、点、千分位逗号，排除已格式化的「2,000.00」
_NUM_RUN = re.compile(r"(?<![\d.,])\d+(?:\s+\d{1,3})*(?![\d.,])")
# 键值表补头：内容列里这些形状算「表头样」——短且无中文标点，或是纯数字
_KV_PUNCT = re.compile(r"[（）【】《》「」、，。；：]")


# ——— PDF 释义表 ———
# 招股书开头的「词语 …… 指 …… 含义」表没有框线，find_tables 靠线条
# 反推单元格，这种表整页漏掉。好在它的版式极其固定：孤立的「指」作为
# 连接词竖着对成一列，左边是词条、右边是释义——拿「指」当锚点，按左右
# 两列拼行就能还原，不需要猜任何表格线。

# 连接词：锚点必须是整个词（「本指导」里的指不算）
_GLOSS_ANCHORS = {"指", "系指", "是指"}
# 一页至少这么多行才当释义表——正文里偶尔也有孤立的「指」，数量撑不起来
_GLOSS_MIN_ROWS = 3
# 锚点列的对齐容差：同一列的「指」x0 相差不会超过这个数
_GLOSS_X_TOL = 18.0


def _pdf_glossary_tables(
    doc, exclude: list[tuple[int, tuple[float, float, float, float]]]
) -> list[tuple[int, float, tuple, list[list[str]]]]:
    """识别无框释义表，返回值结构与 _pdf_tables 相同（可直接合并归并）。"""
    out: list[tuple[int, float, tuple, list[list[str]]]] = []
    for pno in range(doc.page_count):
        try:
            words = doc[pno].get_text("words")
        except Exception:
            continue
        rows, bbox = _glossary_rows(_word_lines(words))
        if len(rows) < _GLOSS_MIN_ROWS or bbox is None:
            continue
        if any(
            pno + 1 == page and _overlap_ratio(bbox, box) > 0.3
            for page, box in exclude
        ):
            # 这片区域 find_tables 已经抽到了（页上还残留些框线），
            # 再出一份就是重复切片
            continue
        out.append(
            (
                pno + 1,
                float(bbox[1]),
                bbox,
                # Markdown 表格首行即表头（前端渲染成灰底），补一列语义
                # 表头，别把第一条释义顶上去当列名
                [["词语", "含义"], *rows],
            )
        )
    return out


def _word_lines(
    words: list[tuple[float, float, float, float, str, int, int, int]],
) -> list[list[tuple[float, float, float, float, str]]]:
    """pymupdf 的 words → 视觉行（按 y0 聚类，行内按 x 排序）。"""
    out: list[list[tuple[float, float, float, float, str]]] = []
    for w in sorted(words, key=lambda w: (w[1], w[0])):
        item = (w[0], w[1], w[2], w[3], w[4])
        if out and item[1] - out[-1][0][1] <= 4:
            out[-1].append(item)
        else:
            out.append([item])
    for line in out:
        line.sort(key=lambda w: w[0])
    return out


def _glossary_rows(
    lines: list[list[tuple[float, float, float, float, str]]],
) -> tuple[list[list[str]], tuple | None]:
    """视觉行 → [词条, 释义] 二维行。

    每条释义由锚点行开头，随后的行只要没有字压在锚点列上（词条换行 /
    释义换行，或两边同时换行）都算这条的延续；一旦某行的字横跨锚点列，
    说明表格结束了、回到正文。
    """
    anchors: list[tuple[int, tuple[float, float, float, float, str]]] = []
    for i, line in enumerate(lines):
        for w in line:
            if w[4] in _GLOSS_ANCHORS:
                anchors.append((i, w))
                break
    if len(anchors) < _GLOSS_MIN_ROWS:
        return [], None
    xs = sorted(a[1][0] for a in anchors)
    if xs[-1] - xs[0] > _GLOSS_X_TOL:
        return [], None

    def cell_text(span: list[tuple]) -> str:
        """一列里的若干词拼回文字：同视觉行的按间距拼，跨行直接连。"""
        parts: list[str] = []
        for line in span:
            text = ""
            prev = None
            for w in line:
                if prev is not None and w[0] - prev[2] > 2.5:
                    # 中文之间的小间隙是排版字距（"集团公司 有"），直连；
                    # 西文之间留真空格
                    gap_cjk = _CJK_RE.search(prev[4][-1]) or _CJK_RE.match(w[4])
                    text += "" if gap_cjk else " "
                text += w[4]
                prev = w
            if text:
                parts.append(text)
        return "".join(parts)

    rows: list[list[str]] = []
    box: list[float] | None = None
    for k, (i, anchor) in enumerate(anchors):
        ax0, ay0, ax1, _ay1, _t = anchor
        stop = anchors[k + 1][0] if k + 1 < len(anchors) else len(lines)
        term_span: list[list[tuple]] = []
        def_span: list[list[tuple]] = []
        ended = True
        for line in lines[i:stop]:
            # 锚点行上锚点自己的字就在锚点列，不算「压线」
            if all(w is anchor or w[2] < ax0 - 2 or w[0] > ax1 + 2 for w in line):
                left = [w for w in line if w[2] < ax0 - 2 and w is not anchor]
                right = [w for w in line if w[0] > ax1 + 2 and w is not anchor]
                if left:
                    term_span.append(left)
                if right:
                    def_span.append(right)
                for w in line:
                    box = box if box is not None else list(w[:4])
                    box[0] = min(box[0], w[0])
                    box[1] = min(box[1], w[1])
                    box[2] = max(box[2], w[2])
                    box[3] = max(box[3], w[3])
            else:
                # 字压在锚点列上 → 表格结束、回到正文。这条本身在压线行
                # 之前已经收完，照常保留；后面的锚点（正文里恰好孤立的
                # 「指」）不能再收割
                ended = False
                break
        term, definition = cell_text(term_span), cell_text(def_span)
        if term and definition:
            rows.append([term, definition])
        if not ended:
            break
    return rows, tuple(box) if box else None


def _overlap_ratio(
    a: tuple[float, float, float, float], b: tuple[float, float, float, float]
) -> float:
    ox = max(0.0, min(a[2], b[2]) - max(a[0], b[0]))
    oy = max(0.0, min(a[3], b[3]) - max(a[1], b[1]))
    return ox * oy / max((a[2] - a[0]) * (a[3] - a[1]), 1.0)


def _header_has_cjk(raw: list[list[str | None]]) -> bool:
    """表头行是否带中文列名。三线表被切掉「资产」列时，表头只剩日期数字。"""
    if not raw:
        return False
    return any(_CJK_RE.search(cell or "") for cell in raw[0])


def _replace_clipped_tables(found: list, alt: list) -> list:
    """用 lines_strict 的结果替换表头没有中文的表（bbox 重叠的那张）。"""
    used: set[int] = set()
    out: list = []
    for table in found:
        raw = table.extract()
        if _header_has_cjk(raw):
            out.append(table)
            continue
        tb = tuple(table.bbox)
        pick = None
        for j, other in enumerate(alt):
            if j in used:
                continue
            if _overlap_ratio(tb, tuple(other.bbox)) < 0.3:
                continue
            if not _header_has_cjk(other.extract()):
                continue
            pick = (j, other)
            break
        if pick is not None:
            used.add(pick[0])
            out.append(pick[1])
        else:
            out.append(table)
    return out


def _find_page_tables(page) -> list:
    """默认按线条抽表；表头被切掉中文列名时改用 lines_strict。

    招股书三线表只有横线和数字区的短竖线，默认策略会把「资产」列和
    最右一列年份切到框外，抽成纯数字网格。lines_strict 认完整外框，
    但不能当默认——合并单元格的表（股东名册、发行费用）列数会变少。
    """
    try:
        found = list(page.find_tables().tables)
    except Exception:
        found = []
    clipped = [t for t in found if not _header_has_cjk(t.extract())]
    if not found or not clipped:
        return found
    try:
        alt = list(page.find_tables(strategy="lines_strict").tables)
    except Exception:
        return found
    if not alt:
        return found
    return _replace_clipped_tables(found, alt)


def _pdf_tables(
    doc,
) -> tuple[list[tuple[int, float, tuple, list[list[str]]]], dict[int, float]]:
    """逐页找表格，返回 (表列表, 页高)。表是 (页码, 顶部纵坐标, bbox, 清洗后的二维行)。"""
    out: list[tuple[int, float, tuple, list[list[str]]]] = []
    heights: dict[int, float] = {}
    for pno in range(doc.page_count):
        try:
            heights[pno + 1] = float(doc[pno].rect.height)
        except Exception:
            heights[pno + 1] = 0.0
        try:
            tables = _find_page_tables(doc[pno])
        except Exception:
            logger.warning("PDF 第 %d 页表格识别失败，跳过", pno + 1, exc_info=True)
            continue
        for table in tables:
            # 一张图里可能拆出多张表（见 _split_repeated_header）：共用同一
            # 个 bbox 即可，正文行照样整块剔掉。外框常比最后一行格子大一截，
            # 会把表下的签字栏（法定代表人）也当成表格内文本删掉。
            bbox = _table_content_bbox(table)
            for rows in _clean_table(table.extract()):
                if len(rows) < 2 or max(len(r) for r in rows) < 2:
                    continue
                out.append((pno + 1, float(bbox[1]), bbox, rows))
    return _carry_page_headers(out, heights), heights


def _carry_page_headers(
    tables: list[tuple[int, float, tuple, list[list[str]]]],
    page_heights: dict[int, float],
) -> list[tuple[int, float, tuple, list[list[str]]]]:
    """跨页续表补表头。

    PDF 里一张表排到页底、下一页接着排时，第二页**不重印列名**，抽出来
    第一行就是数据。原样渲染成 Markdown 后首行被当成表头，于是
    「基本每股收益（元/股） | 0.01 | …」成了列名，后面所有数字都没有
    列名可对——而 `_chunk_table` 又会把这一行当表头复制到每一片。

    补的条件同时成立才动手，宁可漏补也不误补：
    - 表格在页面顶部（续表只会从顶部开始）；
    - 首行是数据行（含纯数字单元格）；
    - 上一页最后一张表排到了页底（说明它是被切断的，不是自己结束的）；
    - 列数与最近一次见到的表头一致。
    """
    prev_bottom: dict[int, float] = {}
    for pno, _y, bbox, _rows in tables:
        prev_bottom[pno] = float(bbox[3])

    out: list[tuple[int, float, tuple, list[list[str]]]] = []
    last_header: list[str] | None = None
    for pno, y, bbox, rows in tables:
        width = max(len(row) for row in rows)
        if (
            y <= _HEAD_CARRY_Y
            and _is_data_row(rows[0])
            and last_header is not None
            and len(last_header) == width
            and _reached_page_bottom(
                prev_bottom.get(pno - 1), page_heights.get(pno - 1)
            )
        ):
            rows = [_tidy_header(last_header), *rows]
        out.append((pno, y, bbox, rows))
        if _is_header_row(rows[0]):
            last_header = rows[0]
    return out


def _reached_page_bottom(bottom: float | None, height: float | None) -> bool:
    if not bottom or not height:
        return False
    return bottom >= height - _PAGE_BOTTOM_SLACK


def _merge_continued_tables(
    tables: list[tuple[int, float, tuple, list[list[str]], str | None]],
    page_heights: dict[int, float],
) -> list[tuple[int, float, tuple, list[list[str]], str | None, int]]:
    """把续表拼回一张。

    断开的常见原因有两种，判据分开写：
    - 跨页：上一张排到页底、下一张从页顶接着（列名可能重印，也可能只剩数据）；
    - 同页：识别器在缺线、注释处把一张表切成上下两框，框距很近、列数相同。

    左右并排的表 y 接近但 x 几乎不重叠，不能并。表题不同也不能并。
    """
    if not tables:
        return []
    ordered = sorted(tables, key=lambda item: (item[0], item[1]))
    out: list[tuple[int, float, tuple, list[list[str]], str | None, int]] = []
    for page, y, bbox, rows, caption in ordered:
        if out and _is_table_continuation(
            out[-1], (page, y, bbox, rows, caption), page_heights
        ):
            p0, y0, b0, r0, c0, _end = out[-1]
            out[-1] = (p0, y0, b0, _concat_table_rows(r0, rows), c0 or caption, page)
        else:
            out.append((page, y, bbox, rows, caption, page))
    return out


def _concat_table_rows(
    head: list[list[str]], tail: list[list[str]]
) -> list[list[str]]:
    if tail and head and tail[0] == head[0]:
        return [*head, *tail[1:]]
    return [*head, *tail]


def _is_table_continuation(
    prev: tuple[int, float, tuple, list[list[str]], str | None, int],
    cur: tuple[int, float, tuple, list[list[str]], str | None],
    page_heights: dict[int, float],
) -> bool:
    p_page, _py, p_bbox, p_rows, p_cap, _pend = prev
    c_page, c_y, c_bbox, c_rows, c_cap = cur
    if not p_rows or not c_rows:
        return False
    if max(len(r) for r in p_rows) != max(len(r) for r in c_rows):
        return False
    if p_cap and c_cap and p_cap != c_cap:
        return False
    # 续表要么重印同一行表头，要么第一行已经是数据
    if p_rows[0] != c_rows[0] and not _is_data_row(c_rows[0]):
        return False
    if not _boxes_vertically_stacked(p_bbox, c_bbox):
        return False
    if c_page == p_page + 1:
        return c_y <= _HEAD_CARRY_Y and _reached_page_bottom(
            float(p_bbox[3]), page_heights.get(p_page)
        )
    if c_page == p_page:
        gap = c_y - float(p_bbox[3])
        return -8.0 <= gap <= _SAME_PAGE_GAP
    return False


def _boxes_vertically_stacked(prev_bbox: tuple, cur_bbox: tuple) -> bool:
    """左右并排的表 x 几乎不重叠，不能当成上下续表。"""
    px0, _py0, px1, _py1 = (float(v) for v in prev_bbox)
    cx0, _cy0, cx1, _cy1 = (float(v) for v in cur_bbox)
    overlap = min(px1, cx1) - max(px0, cx0)
    width = min(px1 - px0, cx1 - cx0)
    return width > 0 and overlap >= width * 0.5


def _tidy_header(row: list[str]) -> list[str]:
    """表头首格常和第一条数据粘在一起（「项目流动资产：」），还原成「项目」。"""
    return [re.sub(r"^项目[^：]{1,10}：$", "项目", cell) if cell else cell for cell in row]


def _attach_captions(
    tables: list[tuple[int, float, tuple, list[list[str]]]],
    lines: list[sections.Line],
) -> list[tuple[int, float, tuple, list[list[str]], str | None]]:
    """给每张表找回表题，返回多带一个 caption 字段的五元组。

    为什么必须做：表格切片的正文只有「项目 | 本期 | 上期」和一堆数字，
    **表题（「母公司利润表」）不进来**。于是"母公司营业收入是多少"这种问题，
    数值所在的切片里根本没有"母公司"三个字——向量和关键词都够不着，只能
    靠相邻的标题切片勉强召回，而那一片没有数值。同理，图谱抽取也拿不到
    "这张表属于谁"，metric 实体因此连不到 org 上（实测孤立率 44%）。

    用 kept（已剔页眉页脚）而不是原始行：页眉每页重复，直接往上找会把
    「XX公司2026年第三季度报告」误当表题。
    """
    out: list[tuple[int, float, tuple, list[list[str]], str | None]] = []
    prev_header: list[str] | None = None
    prev_caption: str | None = None
    for page, y, bbox, rows in tables:
        caption = _caption_above(lines, page, y)
        if (
            caption is None
            and y <= _HEAD_CARRY_Y
            and prev_header is not None
            and rows[0] == prev_header
        ):
            # 跨页续表：表题只在首页印一次。必须同时满足「从页顶开始」和
            # 「表头与上一张表一致」（_carry_page_headers 补过的表头就是上一张的），
            # 否则会把上一节的标题安到一张无关的表上——顺丰第 5 页三张
            # 页顶表就差点被冠上「（二）非经常性损益项目和金额」
            caption = prev_caption
        # 只跟紧邻的上一张表接续：中间隔了一张没表题的表，链就断了，
        # 免得把更早那张表的标题递到后面去
        prev_caption = caption
        prev_header = rows[0]
        out.append((page, y, bbox, rows, caption))
    return out


def _lines_above(
    lines: list[sections.Line], page: int, y0: float
) -> list[sections.Line]:
    """表体上方窗口内的行，从近到远。表题和单位都从这里捡。"""
    candidates = [
        ln
        for ln in lines
        if ln.page == page
        and ln.bbox
        and y0 - _CAPTION_WINDOW <= ln.bbox[3] <= y0 + 1
    ]
    candidates.sort(key=lambda ln: ln.bbox[3], reverse=True)
    return candidates


def _caption_above(lines: list[sections.Line], page: int, y0: float) -> str | None:
    """在表体上方找最近的一条「像表题」的行，找不到返回 None。

    从近往远走、跳过元信息行：正常的报表上方依次是「单位：千元」「编制单位：」
    「2026 年1—9 月」，真正的表题在更上面一行。
    """
    for ln in _lines_above(lines, page, y0):
        text = ln.text.strip()
        if not text or len(text) > _CAPTION_MAX_CHARS:
            continue
        if _CAPTION_NOISE.search(text):
            continue
        if _CAPTION_SENTENCE.search(text):
            continue
        if _CAPTION_SHAPE.search(text):
            return text
    return None


# 报表上方常把单位、币种、审计状态粘在一行：
# 「单位：千元币种：人民币审计类型：未经审计」。编制单位是主体名，太长
# 也容易和表题抢戏，不收。
_UNIT_FIELD = re.compile(
    r"(?:金额单位|(?<![编准])单位)[:：]\s*(.+?)(?=币种|审计类型|编制单位|$)"
)
_CURRENCY_FIELD = re.compile(
    r"币种[:：]\s*(.+?)(?=审计类型|金额单位|(?<![编准])单位|编制单位|$)"
)
_AUDIT_FIELD = re.compile(r"(?:审计类型[:：]\s*)?(未经审计|已审计|经审计)")


def _table_meta_above(lines: list[sections.Line], page: int, y0: float) -> str | None:
    """从表上方捡单位 / 币种 / 审计状态，拼成跟在表题后面的一行。

    这些字不进切片的话，后半张表被切开后只剩数字，模型会把百万元当成元。
    近的行优先：两张表叠在一起时，要的是紧挨着这张表的那行。
    """
    unit = currency = audit = None
    for ln in _lines_above(lines, page, y0):
        text = ln.text.strip()
        if not text:
            continue
        if unit is None:
            m = _UNIT_FIELD.search(text)
            if m:
                unit = m.group(1).strip(" 。；;，,")
        if currency is None:
            m = _CURRENCY_FIELD.search(text)
            if m:
                currency = m.group(1).strip(" 。；;，,")
        if audit is None:
            m = _AUDIT_FIELD.search(text)
            if m:
                audit = m.group(1)
        if unit and currency and audit:
            break
    parts: list[str] = []
    if unit:
        parts.append(unit if unit.startswith("单位") else f"单位：{unit}")
    if currency:
        parts.append(currency if "币种" in currency else f"币种：{currency}")
    if audit:
        parts.append(audit)
    return " · ".join(parts) if parts else None


def _table_prefix(caption: str | None, meta: str | None) -> str:
    """表题 + 单位行。口径判断只用表题，所以两者分开传、拼的时候再合。"""
    return "\n".join(part for part in (caption, meta) if part)


# 报表口径：同一个指标在「合并报表」与「母公司报表」里各有一份数，值不一样，
# 而用户的问法通常不带口径（「中国能建期末现金余额」默认问集团口径）。
# 不区分会错得毫无痕迹：能建期末现金，母公司 2,721,444、合并 77,436,294，
# 相差 28 倍，两片都含「期末现金及现金等价物余额」这串字。
SCOPE_CONSOLIDATED = "consolidated"
SCOPE_PARENT = "parent"
_PARENT_SCOPE = re.compile(r"母公司|个别报表|单体报表")
_CONSOLIDATED_SCOPE = re.compile(r"合并")


def _table_scope(caption: str | None) -> str | None:
    """从表题判断报表口径，判不出返回 None（不猜）。

    只看表题不看表体：表内的「其中：」小计项在两张表里都出现，拿它判口径
    会把合并表和母公司表判成同一个。
    """
    if not caption:
        return None
    if _PARENT_SCOPE.search(caption):
        return SCOPE_PARENT
    if _CONSOLIDATED_SCOPE.search(caption):
        return SCOPE_CONSOLIDATED
    return None


def _is_data_row(row: list[str]) -> bool:
    """数据行：至少有一个纯数字单元格。"""
    return any(cell and _NUM_ONLY.match(cell) for cell in row)


def _is_header_row(row: list[str]) -> bool:
    """列名行：至少两个非空单元格、不含纯数字、带中文。"""
    cells = [cell for cell in row if cell]
    if len(cells) < 2 or _is_data_row(row):
        return False
    return any(_CJK_RE.search(cell) for cell in cells)


def _headerish_row(row: list[str]) -> bool:
    """更像列名行：每格都短、都不含数字和句读。

    比 `_is_header_row` 更严，用来在页内认「第二张表的列名」——那里没有
    重复列名可比，只能看形状。
    """
    cells = [cell for cell in row if cell]
    if len(cells) < 2 or _is_data_row(row):
        return False
    return all(
        len(cell) <= _HEADER_CELL_MAX and not _HEADER_SENTENCE.search(cell)
        for cell in cells
    )


def _clean_table(raw: list[list[str | None]]) -> list[list[list[str]]]:
    """单元格清洗 → 拆开拼在一起的表 → 表头归位 → 删空列。

    返回的是「一张图里可能拆出的多张表」：pymupdf 常把上下排的两个表
    识成一个（第二段的表头变成数据行），导致表头和数据对不上列。
    """
    rows = [[_clean_cell(cell) for cell in row] for row in raw]
    rows = [row for row in rows if any(row)]
    if not rows:
        return []
    width = max(len(row) for row in rows)
    rows = [row + [""] * (width - len(row)) for row in rows]
    out: list[list[list[str]]] = []
    for part in _split_repeated_header(rows):
        # 先删空列：合并单元格常切出 6 列网格，表头在 1/4、数据在 0/3，
        # 去空后才变成 split_interleaved 认得的隔列交错
        part = _drop_empty_cols(part)
        part = _merge_header(part)
        part = _split_interleaved(part)
        part = _align_header(part)
        out.append(_add_kv_header(_fill_header_corner(_drop_empty_cols(part))))
    return out


def _fill_header_corner(rows: list[list[str]]) -> list[list[str]]:
    """表头左上角空着就补「项目」。

    合并单元格切出来的表头常是「空 | 本报告期末 | 上年度末」——Markdown
    首行即列名，空着的那格就成了无名列。
    """
    if len(rows) < 2 or len(rows[0]) < 3 or rows[0][0]:
        return rows
    filled = sum(1 for row in rows[1:] if row[0])
    if filled * 2 < len(rows) - 1:
        return rows
    return [["项目", *rows[0][1:]], *rows[1:]]


def _add_kv_header(rows: list[list[str]]) -> list[list[str]]:
    """两列键值表（"股票类型 | 人民币普通股（A 股）"）没有表头行。

    Markdown 表格首行必被渲染成表头（灰底），不补的话第一条发行概况就
    被顶成了列名。判据：每行都是「短标签 + 内容」，且没有一行的内容列
    长得像表头（短而无中文标点、或纯数字）——后者说明首行本来就是列名。
    """
    if len(rows) < 2 or any(len(r) != 2 for r in rows):
        return rows
    if not all(r[0] and r[1] and len(r[0]) <= 10 for r in rows):
        return rows
    if any(_kv_headerish(r[1]) for r in rows):
        return rows
    return [["项目", "内容"], *rows]


def _kv_headerish(cell: str) -> bool:
    """内容列是否长得像表头：「数值」「含义」这类一两个字，或纯数字。

    刻意只认 ≤2 字——「余额包销」这种 4 字短内容是真数据；代价是
    「解释说明」这类 4 字表头认不出来，但那比把数据行当表头少见得多。
    """
    return _NUM_ONLY.match(cell) is not None or (
        len(cell) <= 2 and not _KV_PUNCT.search(cell)
    )


def _split_repeated_header(rows: list[list[str]]) -> list[list[list[str]]]:
    """body 里又出现表头行 → 从那里断成两张表。

    两条判据任一成立即可，但都要求前面已有表体、前面那段表体多数行带
    数字（数据表才有；释义是整句）：

    - **重印列名**：非空列与首行完全一致，且有 ≥2 字同名的列名。这是
      「同一张表翻页后重印表头」；
    - **页内叠表**：这一行长得像列名（每格都短、无数字、无句读），且紧
      挨着的上一行是数据行。这是「上下排的两张表被识成一张」，第二张的
      列名接在第一张的数据后面，列占用常与首行不同（合并单元格切出的
      空列位置不一样），所以不能要求列占用一致。

    释义表（"X 指 Y"）每条的行形状都一样，只靠「重印列名」会被逐条切开，
    一张表碎成十几张两行小表——所以那条还要求 `_same_head_label`。
    """
    if len(rows) < 3:
        return [rows]
    head_cols = {i for i, cell in enumerate(rows[0]) if cell}
    parts: list[list[list[str]]] = []
    cur: list[list[str]] = [rows[0]]
    for row in rows[1:]:
        cols = {i for i, cell in enumerate(row) if cell}
        numeric = any(cell and _NUM_ONLY.match(cell) for cell in row)
        reprinted = cols == head_cols and _same_head_label(rows[0], row)
        stacked = _headerish_row(row) and any(_is_data_row(r) for r in cur[-2:])
        if (reprinted or stacked) and len(cur) > 1 and not numeric and _mostly_numeric(cur[1:]):
            parts.append(cur)
            cur = [row]
        else:
            cur.append(row)
    parts.append(cur)
    return parts


def _same_head_label(head: list[str], row: list[str]) -> bool:
    """重复出现的表头得跟首行有共同列名（≥2 字，排掉「指」这类连接词）。"""
    pairs = zip(head, row)
    return any(len(a) >= 2 and a == b for a, b in pairs)


def _mostly_numeric(body: list[list[str]]) -> bool:
    """表体多数行带数字才像数据表——用来把数据表和整句的释义区分开。"""
    if not body:
        return False
    hits = sum(1 for row in body if any(re.search(r"\d", cell) for cell in row))
    return hits * 2 >= len(body)


def _fill_wrap_punct(text: str, marks: list[str]) -> str:
    """把单独成行的中文标点填进上一行短语之间的空格，多出来的贴行尾。

    「负责行业研究 收集…投资 合作机会」+「， 、 ，」→
    「负责行业研究，收集…投资、合作机会」。空格比标点多时，剩下的
    空格留给后面的 _CJK_SPACE 去拼（行尾被切断的「新项」「目」）。
    """
    if not marks:
        return text
    parts = re.split(r"(\s+)", text)
    i = 0
    for k, part in enumerate(parts):
        if part.isspace() and i < len(marks):
            parts[k] = marks[i]
            i += 1
    return "".join(parts) + "".join(marks[i:])


def _clean_cell(value: str | None) -> str:
    text = value or ""
    if "\n" in text:
        # 多行单元格里的孤儿标点行分三类（都是 PDF 把标点排成独立块）：
        # - 只有点的「. / , .」：是前面数字被甩出的小数点（可带千分位逗号），
        #   按顺序补回数字里——「365元/股\n.\n…」是 3.65、「200000万元」
        #   是 2,000.00；
        # - 纯中文标点（「， 、 ；」）：插回上一行空格缝，职责表靠这个还原；
        # - 中英文混杂（", ，" / ". 。 ，"）：数字碎片和句读搅在一起，
        #   位置不可考，删掉。末行句号仍保留。
        dots = 0
        kept: list[str] = []
        lines = text.split("\n")
        for i, ln in enumerate(lines):
            if _DOT_ROW.match(ln.strip()):
                dots += 1
            elif _WRAP_PUNCT.match(ln) and kept:
                kept[-1] = _fill_wrap_punct(
                    kept[-1], _WRAP_PUNCT_MARK.findall(ln)
                )
            elif _PUNCT_ROW.match(ln) and i < len(lines) - 1:
                continue
            else:
                kept.append(ln)
        text = "\n".join(kept)
        if dots:
            text = _restore_dots(text, dots)
    text = text.replace("\n", " ")
    text = _CELL_WS.sub(" ", text).strip()
    # 先修连写的日期（"20090807 . ."），否则会被下面的数字重组当成
    # 「整数连写 + 甩到末尾的小数点」切坏
    text = _DATE_DOTS.sub(r"\1.\2.\3", text)
    text = _DATE_DOTS_MD.sub(r"\1.\2.\3", text)
    text = _regroup_tail(text)
    m = _NUM_THOUSANDS.match(text)
    if m:
        groups = m.group(1).split()
        text = ",".join(groups[:-1]) + "." + groups[-1] + m.group(2)
    else:
        text = _NUM_DECIMAL.sub(
            lambda m: f"{m.group(1)}.{m.group(2)}{m.group(3)}", text
        )
        text = _NUM_RANGE.sub(r"\1-\2", text)
    text = _CJK_SPACE.sub(r"\1\2", text)
    # 数字与中文之间也是排版间距："2005年12月31 日" → "2005年12月31日"
    text = re.sub(rf"([{_CJK}])\s+(\d)", r"\1\2", text)
    text = re.sub(rf"(\d)\s+([{_CJK}])", r"\1\2", text)
    text = _SPACE_BEFORE_PUNCT.sub(r"\1", text)
    text = _SPACE_AFTER_PUNCT.sub(r"\1", text)
    text = _MULTI_COMMA.sub("，", text)
    text = _DATE_DASH.sub(r"\1-\2-\3", text)
    # 日期区间里的空格："2009.08.07- 2019.08.06"
    text = re.sub(r"(\d)\s*-\s+(\d)", r"\1-\2", text)
    m = _NUM_TRAIL_SIGNS.match(text)
    if m:
        text = ("-" if m.group(2) else "") + m.group(1)
    text = _SPACE_AFTER_PUNCT_2.sub(r"\1", text)
    text = _TRAIL_COMMA.sub("", text)
    text = _CJK_COMMA.sub("", text)
    # 表格里也有勾选（「是否 √是 □否」），和正文同一个符号字体问题
    return _escape_cell(normalize_choices(text))


def _regroup_tail(text: str) -> str:
    """把甩到单元格末尾的小数点/逗号拼回数字里。"""
    for pattern, decimal in (
        (_NUM_TAIL_BOTH, True),
        (_NUM_TAIL_DOT, True),
        (_NUM_TAIL_COMMA, False),
    ):
        m = pattern.match(text)
        if not m:
            continue
        prefix, digits, suffix = m.group(1), m.group(2), m.group(3)
        if decimal:
            if len(digits) > 2:
                return f"{prefix}{_thousands(digits[:-2])}.{digits[-2:]}{suffix}"
            return f"{prefix}0.{digits}{suffix}"
        return f"{prefix}{_thousands(digits)}{suffix}"
    return text


def _thousands(digits: str) -> str:
    """4018 → 4,018（三位一组，从右往左）"""
    if len(digits) <= 3:
        return digits
    out: list[str] = []
    while len(digits) > 3:
        out.insert(0, digits[-3:])
        digits = digits[:-3]
    out.insert(0, digits)
    return ",".join(out)


def _restore_dots(text: str, count: int) -> str:
    """把甩成孤儿行的小数点按出现顺序补回数字（招股书惯例：两位小数）。

    「365元/股」→ 3.65 元/股、「200000万元」→ 2,000.00 万元；多段连写
    （"5176 63"）末段是小数 → 5,176.63 同 _NUM_THOUSANDS 的拼法。只动前
    count 处：日期（2016年12月31日）排在数字主体后面，不会误伤。
    """
    remaining = count

    def sub(m: re.Match[str]) -> str:
        nonlocal remaining
        parts = m.group(0).split()
        if len(parts) == 1 and len(parts[0]) < 2:
            return m.group(0)  # 单个一位数字不可能是「整数+甩点」，不消耗
        if remaining <= 0:
            return m.group(0)
        remaining -= 1
        if len(parts) > 1:
            return ",".join(parts[:-1]) + "." + parts[-1]
        digits = parts[0]
        if len(digits) > 2:
            return f"{_thousands(digits[:-2])}.{digits[-2:]}"
        return f"0.{digits}"

    return _NUM_RUN.sub(sub, text)


def _align_header(rows: list[list[str]]) -> list[list[str]]:
    """表头归位：合并单元格常把表头挤到右边一列，和数据对不上。

    「销量（吨）」的格子横跨两列时，网格切出来的表头落在 col1，而数值
    落在 col1/col2 —— 直接渲染的话每列都串了位。这里按「有数据的列」
    依次安放表头文字：数量对得上才动，对不上说明不是这种表，不猜。
    """
    if len(rows) < 2:
        return rows
    header, body = rows[0], rows[1:]
    data_cols = [i for i in range(len(header)) if any(row[i] for row in body)]
    fills = [cell for cell in header if cell]
    if not fills or len(fills) != len(data_cols):
        return rows
    aligned = [""] * len(header)
    for col, cell in zip(data_cols, fills):
        aligned[col] = cell
    return [aligned] + body


def _merge_header(rows: list[list[str]]) -> list[list[str]]:
    """表头占了多行时拼回一行，最多吃三行（表头再复杂就认命）。

    三种形态，判据是「第二行的非空列和第一行重不重叠」：

    - **续行**：第二行的列是第一行的子集，且都是短碎片——"销售收入（万"
      + "元）"、"占总收入比例" + "（%"）。长内容说明那是数据行，不动；
    - **分散**：两行的非空列完全不重叠——招股书常见的斜线表头把「年份」
      放第 0 列、「2006年1-6月」放第 1~4 列，分了两三行才写完。
    - **分组**：第一行是合并单元格的分组名（「中国银行集团」「中国银行」），
      第二行填期间 / 审计状态。两行部分重叠，按分组名向右沿用再拼到每一列。

    部分重叠但第二行已经是数字，一律不动——宁可留着多行表头，
    也不能把两行数据粘成一行。
    """
    for _ in range(3):
        if len(rows) < 2:
            break
        first, second = rows[0], rows[1]
        f_cols = {i for i, cell in enumerate(first) if cell}
        s_cols = {i for i, cell in enumerate(second) if cell}
        if not f_cols or not s_cols:
            break
        if s_cols <= f_cols:
            cells = [second[i] for i in s_cols]
            if any(len(cell) > _MAX_HEADER_CELL for cell in cells):
                break
            # 纯数字的一律是数据行（"100"、"200"），拼进表头就串味了
            if any(_NUM_ONLY.match(cell) for cell in cells):
                break
            merged = _concat_header_rows(first, second)
        elif s_cols & f_cols:
            merged = _flatten_group_header(first, second)
            if merged is None:
                break
        elif _alternating(f_cols, s_cols):
            # 两行的非空列均匀交替（表头在 1/4/7 列、数据在 0/3/6 列）不是
            # 分两行写的表头，而是合并单元格把表头和数据挤进了同一套网格。
            # 合并起来就成了「1 | 序号 | 孙瑞良 | 股东名称」这种串味的行，
            # 交给后面的 _align_header 按数据列归位
            break
        else:
            merged = _concat_header_rows(first, second)
        rows = [merged] + rows[2:]
    return rows


def _concat_header_rows(first: list[str], second: list[str]) -> list[str]:
    merged = list(first)
    for i, cell in enumerate(second):
        if not cell:
            continue
        merged[i] = merged[i] + cell if merged[i] else cell
    return merged


def _flatten_group_header(first: list[str], second: list[str]) -> list[str] | None:
    """把「分组名跨列 + 下面一行期间」摊成每列一个完整列名。

    只在第二行像列名、不像数据时动手。股东名册那种表头/数据交错
    走的是另一条路，不会进这里（非空列互不重叠）。
    """
    if _is_data_row(second) or not _headerish_row(second):
        return None
    filled = _carry_span_labels(first)
    return [_join_header_cell(a, b) for a, b in zip(filled, second)]


def _carry_span_labels(row: list[str]) -> list[str]:
    """合并单元格留下的空列，用左侧分组名填上。左上角「项目」不往右传。"""
    out = list(row)
    corner = out[0] if out else ""
    last = ""
    for i, cell in enumerate(out):
        if cell:
            last = cell
            continue
        if i > 0 and last and last != corner:
            out[i] = last
    return out


def _join_header_cell(top: str, bottom: str) -> str:
    if not bottom:
        return top
    if not top or bottom.startswith(top) or top in bottom:
        return bottom
    if top.startswith(bottom):
        return top
    return top + bottom


def _split_interleaved(rows: list[list[str]]) -> list[list[str]]:
    """表头被并进数据行：拆回「一行表头 + 纯数据行」。

    跨页续表、合并单元格常抽成交错列 [数据1, 表头1, 数据2, 表头2 …]：
    表头只在首行出现、后续行对应位置留空，直接渲染就是「1 | 序号 | 孙瑞良
    | 股东名称」这种串味的行。

    判据刻意很严，宁可不动也不能把真的宽表拆坏：列数成对；表头位有中文
    且不是纯数字；其余行在表头位要么空、要么与首行一致（不一致说明那是
    真数据列）。合计行例外：标签常写在表头列上，不能因此放弃拆分。
    """
    if len(rows) < 2:
        return rows
    width = len(rows[0])
    if width < 4 or width % 2:
        return rows
    for start in (1, 0):
        head = rows[0][start::2]
        if sum(1 for cell in head if cell) < 2:
            continue
        if not any(_CJK_RE.search(cell) for cell in head):
            continue
        if any(cell and _NUM_ONLY.match(cell) for cell in head):
            continue
        if any(
            not _header_slot_ok(row[start::2], head) for row in rows[1:]
        ):
            continue
        data_start = 1 - start
        body = [_data_slots(row, data_start, start) for row in rows[1:]]
        if not any(any(row) for row in body):
            continue
        return _drop_empty_cols([list(head)] + body)
    return rows


def _header_slot_ok(slots: list[str], head: list[str]) -> bool:
    """表头位要么空/与首行相同，要么是合计行把标签写在了表头列。"""
    if _TOTAL_LABEL.search("".join(slots)):
        return True
    return all(not cell or cell == h for cell, h in zip(slots, head))


def _data_slots(row: list[str], data_start: int, head_start: int) -> list[str]:
    data = list(row[data_start::2])
    if any(data):
        return data
    heads = list(row[head_start::2])
    return heads if any(heads) else data


def _alternating(cols_a: set[int], cols_b: set[int]) -> bool:
    """两组列号是否均匀交替（a 在 1/4/7，b 在 0/3/6）。"""
    # 两列表被切成 4 列交错时各只有 2 个非空列，3 会漏掉（南岭发行费用表）
    if len(cols_a) < 2 or len(cols_b) < 2:
        return False
    seq = sorted(cols_a | cols_b)
    if len(seq) != len(cols_a) + len(cols_b):
        return False
    flags = [i in cols_a for i in seq]
    return all(x != y for x, y in zip(flags, flags[1:]))


def _drop_empty_cols(rows: list[list[str]]) -> list[list[str]]:
    """删掉整列为空的列：跨列合并的单元格会留下一串空列。"""
    width = max(len(row) for row in rows)
    keep = [i for i in range(width) if any(row[i] for row in rows)]
    return [[row[i] for i in keep] for row in rows]


def _table_content_bbox(table) -> tuple:
    """用最后一行格子收紧表格外框，避免把表下的签字栏包进去。"""
    outer = tuple(table.bbox)
    try:
        rows = table.rows
    except Exception:
        return outer
    bottoms = [
        float(row.bbox[3])
        for row in (rows or [])
        if getattr(row, "bbox", None)
    ]
    if not bottoms:
        return outer
    return (outer[0], outer[1], outer[2], min(float(outer[3]), max(bottoms) + 6.0))


# 签字人姓名：葛海蛟 / 张辉，印在职务上一行，和职务一起留
_SIGNATORY_NAME = re.compile(rf"^[{_CJK}]{{2,4}}$")


def _signatory_pages(lines: list[sections.Line]) -> dict[int, list[float]]:
    pages: dict[int, list[float]] = {}
    for ln in lines:
        if sections.is_signatory(ln):
            pages.setdefault(ln.page, []).append(ln.y_ratio)
    return pages


def _keep_signatory_block(ln: sections.Line, pages: dict[int, list[float]]) -> bool:
    if sections.is_signatory(ln):
        return True
    ys = pages.get(ln.page)
    if not ys or not _SIGNATORY_NAME.match(ln.norm):
        return False
    return any(abs(ln.y_ratio - y) <= 0.08 for y in ys)


def _drop_table_lines(
    lines: list[sections.Line],
    tables: list[tuple[int, float, tuple, list[list[str]]]],
) -> list[sections.Line]:
    """剔掉落在表格框内的文本行：表格已单独成段，正文里不该再有一份。"""
    if not tables:
        return lines
    boxes = [(page, bbox) for page, _y, bbox, _rows in tables]
    sign_pages = _signatory_pages(lines)
    kept: list[sections.Line] = []
    for ln in lines:
        if ln.bbox and _hit_any_box(ln, boxes) and not _keep_signatory_block(ln, sign_pages):
            continue
        kept.append(ln)
    return kept


def _hit_any_box(
    ln: sections.Line, boxes: list[tuple[int, tuple[float, float, float, float]]]
) -> bool:
    """行框与表格框的重叠面积占行面积一半以上，就算这行属于表格。"""
    x0, y0, x1, y1 = ln.bbox
    area = max((x1 - x0) * (y1 - y0), 1.0)
    for page, (bx0, by0, bx1, by1) in boxes:
        if page != ln.page:
            continue
        ox = max(0.0, min(x1, bx1) - max(x0, bx0))
        oy = max(0.0, min(y1, by1) - max(y0, by0))
        if ox * oy / area >= 0.5:
            return True
    return False


def _page_char_counts(lines: list[sections.Line]) -> list[int]:
    """每页原始字符数（含页眉）：图表页的三道筛按整页文字量判定。"""
    counts: dict[int, int] = {}
    for ln in lines:
        counts[ln.page] = counts.get(ln.page, 0) + len(ln.text)
    if not counts:
        return []
    return [counts.get(page, 0) for page in range(1, max(counts) + 1)]


def _pdf_text_fallback(path: Path) -> tuple[list[Segment], list[int]]:
    try:
        from pypdf import PdfReader
    except ImportError as e:
        raise ParseError("服务端缺少 pypdf 依赖，请联系管理员") from e
    try:
        reader = PdfReader(str(path))
    except Exception as e:
        raise ParseError(f"PDF 文件无法打开（{type(e).__name__}）") from e

    segments: list[Segment] = []
    text_lens: list[int] = []
    for i, page in enumerate(reader.pages):
        try:
            text = (page.extract_text() or "").strip()
        except Exception:
            text = ""
        text_lens.append(len(text))
        if text:
            segments.append(Segment(normalize_choices(text), i + 1, TEXT))
    return segments, text_lens


def _describe_parallel(jobs: list[tuple[bytes, str]]) -> list[str | None]:
    """并行送视觉模型，结果顺序与 jobs 一致。

    慢的是视觉调用（每次数秒），页面渲染/图片提取是本地快的——所以先把
     material 准备好，再按 kb_vision_concurrency 分批并行调用。
    熔断后 describe_image 直接返回 None，剩余任务秒过，不会白等。
    """
    workers = max(1, min(settings.kb_vision_concurrency, len(jobs)))
    if workers <= 1:
        return [vision.describe_image(blob, mime) for blob, mime in jobs]
    with ThreadPoolExecutor(
        max_workers=workers, thread_name_prefix="kb-vision"
    ) as pool:
        return list(
            pool.map(lambda job: vision.describe_image(job[0], job[1]), jobs)
        )


def _pdf_figures(path: Path, text_lens: list[int] | None = None) -> list[Segment]:
    """含图页渲染成 PNG 交视觉模型描述；关闭或失败都只是少几段，不阻断解析。"""
    if not vision.available():
        return []
    try:
        import pymupdf
    except ImportError:
        logger.warning("未安装 pymupdf，跳过 PDF 图表识别")
        return []

    out: list[Segment] = []
    try:
        doc = pymupdf.open(str(path))
    except Exception:
        logger.warning("pymupdf 打开 PDF 失败，跳过图表识别", exc_info=True)
        return []
    # 先扫出预算内的含图页（不渲染），再分批「渲染 → 并行识别」；
    # 熔断检查放在批间：熔断后后续页不再渲染（与串行版语义一致）
    try:
        targets: list[int] = []
        for i, page in enumerate(doc):
            text_len = text_lens[i] if text_lens and i < len(text_lens) else None
            if not _page_has_chart(page, text_len):
                continue
            if len(targets) >= settings.kb_vision_max_pages:
                logger.info(
                    "kb 图表识别达到上限 %d 页，其余页跳过",
                    settings.kb_vision_max_pages,
                )
                break
            targets.append(i)
        for start in range(0, len(targets), settings.kb_vision_concurrency):
            if not vision.available():
                break
            pngs: list[tuple[int, bytes]] = []
            for i in targets[start : start + settings.kb_vision_concurrency]:
                try:
                    pngs.append(
                        (
                            i,
                            doc[i]
                            .get_pixmap(dpi=settings.kb_vision_dpi)
                            .tobytes("png"),
                        )
                    )
                except Exception:
                    logger.warning("PDF 第 %d 页渲染失败，跳过", i + 1, exc_info=True)
            if not pngs:
                continue
            descs = _describe_parallel([(png, "image/png") for _i, png in pngs])
            for (i, _png), desc in zip(pngs, descs):
                if desc:
                    out.append(Segment(desc, i + 1, FIGURE))
    finally:
        doc.close()
    return out


def _page_has_chart(page, text_len: int | None = None) -> bool:
    """判断这一页值不值得送视觉模型。

    三道筛子，从便宜到贵：

    1. **文字密度**：正文超过 kb_vision_max_text_chars 的页面直接排除。
       招股书/财报里整页表格会画出几百条矢量线，按「路径数 ≥ 150」判断会
       把表格页统统误认成图表（实测 331 页的招股书误判 124 页，预算 20 页
       全烧在目录和财务报表上）；
    2. **位图面积**：只有占页面面积 ≥ kb_vision_min_image_ratio 的图才算数，
       页眉 logo、页脚装饰线、水印都远小于这条线；
    3. **矢量路径**：ECharts 走 SVG 导出时没有位图，退化成看路径数量，
       但仍要求该页文字很少——纯矢量图表页本来就没多少可抽的文字。
    """
    if text_len is not None and text_len > settings.kb_vision_max_text_chars:
        return False
    try:
        if _has_large_image(page):
            return True
        return len(page.get_drawings()) >= settings.kb_vision_vector_paths
    except Exception:
        return False


def _has_large_image(page) -> bool:
    """页面上是否存在占面积够大的位图（按 bbox 算，不看像素总数）。"""
    try:
        rect = page.rect
        page_area = max(rect.width * rect.height, 1.0)
        for info in page.get_image_info():
            x0, y0, x1, y1 = info["bbox"]
            if (x1 - x0) * (y1 - y0) / page_area >= settings.kb_vision_min_image_ratio:
                return True
    except Exception:
        # get_image_info 拿不到就退回「有没有图」的粗判
        return bool(page.get_images(full=True))
    return False


# ——— DOCX ———


def _parse_docx(path: Path) -> list[Segment]:
    try:
        from docx import Document
        from docx.table import Table
        from docx.text.paragraph import Paragraph
    except ImportError as e:
        raise ParseError("服务端缺少 python-docx 依赖，请联系管理员") from e
    try:
        doc = Document(str(path))
    except Exception as e:
        raise ParseError(f"DOCX 文件无法打开（{type(e).__name__}）") from e

    segments: list[Segment] = []
    buffer: list[str] = []

    def flush() -> None:
        if buffer:
            text = "\n\n".join(buffer).strip()
            if text:
                segments.append(Segment(text, None, TEXT))
            buffer.clear()

    # 按文档顺序遍历段落与表格：表格位置决定上下文，不能统统挪到文末
    for item in doc.iter_inner_content():
        if isinstance(item, Paragraph):
            line = item.text.strip()
            if line:
                buffer.append(line)
        elif isinstance(item, Table):
            markdown = _table_to_markdown(item)
            if markdown:
                flush()
                segments.append(Segment(markdown, None, TABLE))
    flush()

    segments.extend(_docx_figures(doc))
    segments.extend(_docx_embedded_sheets(doc))

    if not segments:
        raise ParseError("DOCX 中没有文本内容")
    return segments


def _rows_to_markdown(rows: list[list[str]]) -> str:
    """二维单元格 → Markdown 表格。PDF/DOCX/Excel 三路共用同一份输出格式。"""
    if not rows:
        return ""
    width = max(len(row) for row in rows)
    rows = [row + [""] * (width - len(row)) for row in rows]
    lines = ["| " + " | ".join(rows[0]) + " |", "|" + "---|" * width]
    lines.extend("| " + " | ".join(row) + " |" for row in rows[1:])
    return "\n".join(lines)


def _table_to_markdown(table) -> str:
    rows: list[list[str]] = []
    for row in table.rows:
        cells = [_escape_cell(cell.text) for cell in row.cells]
        if any(cells):
            rows.append(cells)
    return _rows_to_markdown(rows)


def _docx_figures(doc) -> list[Segment]:
    """Word 里的图片（含粘贴进来的 ECharts 截图）：交给视觉模型转描述。"""
    if not vision.available():
        return []
    try:
        shapes = doc.inline_shapes
        total = len(shapes)
    except Exception:
        return []
    if total > settings.kb_vision_max_pages:
        logger.info(
            "kb 图表识别达到上限 %d 张，其余跳过", settings.kb_vision_max_pages
        )
        total = settings.kb_vision_max_pages
    # 先把图片二进制全部提出来（本地快），再并行送视觉模型
    jobs: list[tuple[bytes, str]] = []
    for i in range(total):
        try:
            blip = shapes[i]._inline.graphic.graphicData.pic.blipFill.blip
            part = doc.part.related_parts[blip.embed]
        except Exception:
            continue
        jobs.append((part.blob, getattr(part, "content_type", "")))
    descs = _describe_parallel(jobs)
    return [Segment(desc, None, FIGURE) for desc in descs if desc]


def _docx_embedded_sheets(doc) -> list[Segment]:
    """Word 里以对象形式嵌入的 Excel：docx 是 zip，内嵌工作簿就在
    word/embeddings/ 下，直接取二进制当独立表格解析。"""
    out: list[Segment] = []
    try:
        rels = list(doc.part.rels.values())
    except Exception:
        return []
    for rel in rels:
        name = (getattr(rel, "target_ref", "") or "").rsplit("/", 1)[-1].lower()
        if not name.endswith((".xlsx", ".xlsm")):
            continue
        try:
            blob = rel.target_part.blob
        except Exception:
            continue
        try:
            out.extend(_parse_xlsx(blob))
        except ParseError:
            logger.warning("DOCX 内嵌 Excel 解析失败：%s", name, exc_info=True)
    return out


# ——— XLSX ———


def _parse_xlsx(source: bytes | Path) -> list[Segment]:
    """每个工作表转一个 Markdown 表格段，page 记 sheet 序号（展示为「表N」）。"""
    try:
        from openpyxl import load_workbook
    except ImportError as e:
        raise ParseError("服务端缺少 openpyxl 依赖，请联系管理员") from e
    # 路径直接交给 openpyxl（read_only 下流式读，超大表不必整体进内存）
    stream = source if isinstance(source, Path) else BytesIO(source)
    try:
        workbook = load_workbook(stream, data_only=True, read_only=True)
    except Exception as e:
        raise ParseError(f"Excel 文件无法打开（{type(e).__name__}）") from e

    segments: list[Segment] = []
    try:
        for index, sheet in enumerate(workbook.worksheets, start=1):
            markdown = _sheet_to_markdown(sheet)
            if markdown:
                segments.append(
                    Segment(f"工作表：{sheet.title}\n{markdown}", index, TABLE)
                )
    finally:
        workbook.close()
    if not segments:
        raise ParseError("Excel 中没有可提取的数据")
    return segments


def _sheet_to_markdown(sheet) -> str:
    rows: list[list[str]] = []
    for row in sheet.iter_rows(values_only=True):
        cells = [_cell_text(value) for value in row]
        if not any(cells):
            continue
        rows.append(cells)
        if len(rows) >= settings.kb_sheet_max_rows:
            logger.info("工作表 %s 超过 %d 行，截断", sheet.title, len(rows))
            break
    if not rows:
        return ""
    width = max(len(row) for row in rows)
    rows = [row + [""] * (width - len(row)) for row in rows]
    # 空表头补列名：Markdown 表格没有列名就无法读懂每列是什么
    rows[0] = [cell or f"列{i + 1}" for i, cell in enumerate(rows[0])]
    return _rows_to_markdown(rows)


def _cell_text(value) -> str:
    if value is None:
        return ""
    if isinstance(value, float) and value.is_integer():
        return str(int(value))
    return _escape_cell(str(value))


def _escape_cell(text: str) -> str:
    # 单元格里的竖线/换行会破坏 Markdown 表格结构
    return " ".join(text.split()).replace("|", "\\|")
