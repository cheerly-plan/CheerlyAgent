"""DOCX / XLSX 解析：Office 文档转带类型的文本段。

从 parsing.py 拆出。Docx 按段落+表格顺序遍历（表格位置决定上下文），
内嵌图片走视觉模型、内嵌 Excel 直接取二进制复用 xlsx 解析；Xlsx 每个
工作表转一个 Markdown 表格段。Markdown 输出格式与 PDF 表格共用
_rows_to_markdown（从 parsing 取，保持单一实现）。
"""

from __future__ import annotations

from io import BytesIO
from pathlib import Path

from ai_agent.core.config import settings
from ai_agent.core.logging import get_logger
from ai_agent.kb.figures import _describe_parallel
from ai_agent.kb import vision
from ai_agent.kb.parsing import (
    FIGURE,
    TABLE,
    TEXT,
    ParseError,
    Segment,
    _escape_cell,
    _rows_to_markdown,
)

logger = get_logger(__name__)


def _parse_docx(path: Path) -> list[Segment]:
    try:
        from docx import Document
        from docx.table import Table
        from docx.text.paragraph import Paragraph
    except ImportError as e:
        raise ParseError("服务端缺少 python-docx 依赖，请联系管理员") from e
    try:
        doc = Document(str(path))
    except Exception as e:
        raise ParseError(f"DOCX 文件无法打开（{type(e).__name__}）") from e

    segments: list[Segment] = []
    buffer: list[str] = []

    def flush() -> None:
        if buffer:
            text = "\n\n".join(buffer).strip()
            if text:
                segments.append(Segment(text, None, TEXT))
            buffer.clear()

    # 按文档顺序遍历段落与表格：表格位置决定上下文，不能统统挪到文末
    for item in doc.iter_inner_content():
        if isinstance(item, Paragraph):
            line = item.text.strip()
            if line:
                buffer.append(line)
        elif isinstance(item, Table):
            markdown = _table_to_markdown(item)
            if markdown:
                flush()
                segments.append(Segment(markdown, None, TABLE))
    flush()

    segments.extend(_docx_figures(doc))
    segments.extend(_docx_embedded_sheets(doc))

    if not segments:
        raise ParseError("DOCX 中没有文本内容")
    return segments


def _table_to_markdown(table) -> str:
    rows: list[list[str]] = []
    for row in table.rows:
        cells = [_escape_cell(cell.text) for cell in row.cells]
        if any(cells):
            rows.append(cells)
    return _rows_to_markdown(rows)


def _docx_figures(doc) -> list[Segment]:
    """Word 里的图片（含粘贴进来的 ECharts 截图）：交给视觉模型转描述。"""
    if not vision.available():
        return []
    try:
        shapes = doc.inline_shapes
        total = len(shapes)
    except Exception:
        return []
    if total > settings.kb_vision_max_pages:
        logger.info(
            "kb 图表识别达到上限 %d 张，其余跳过", settings.kb_vision_max_pages
        )
        total = settings.kb_vision_max_pages
    # 先把图片二进制全部提出来（本地快），再并行送视觉模型
    jobs: list[tuple[bytes, str]] = []
    for i in range(total):
        try:
            blip = shapes[i]._inline.graphic.graphicData.pic.blipFill.blip
            part = doc.part.related_parts[blip.embed]
        except Exception:
            continue
        jobs.append((part.blob, getattr(part, "content_type", "")))
    descs = _describe_parallel(jobs)
    return [Segment(desc, None, FIGURE) for desc in descs if desc]


def _docx_embedded_sheets(doc) -> list[Segment]:
    """Word 里以对象形式嵌入的 Excel：docx 是 zip，内嵌工作簿就在
    word/embeddings/ 下，直接取二进制当独立表格解析。"""
    out: list[Segment] = []
    try:
        rels = list(doc.part.rels.values())
    except Exception:
        return []
    for rel in rels:
        name = (getattr(rel, "target_ref", "") or "").rsplit("/", 1)[-1].lower()
        if not name.endswith((".xlsx", ".xlsm")):
            continue
        try:
            blob = rel.target_part.blob
        except Exception:
            continue
        try:
            out.extend(_parse_xlsx(blob))
        except ParseError:
            logger.warning("DOCX 内嵌 Excel 解析失败：%s", name, exc_info=True)
    return out


def _parse_xlsx(source: bytes | Path) -> list[Segment]:
    """每个工作表转一个 Markdown 表格段，page 记 sheet 序号（展示为「表N」）。"""
    try:
        from openpyxl import load_workbook
    except ImportError as e:
        raise ParseError("服务端缺少 openpyxl 依赖，请联系管理员") from e
    # 路径直接交给 openpyxl（read_only 下流式读，超大表不必整体进内存）
    stream = source if isinstance(source, Path) else BytesIO(source)
    try:
        workbook = load_workbook(stream, data_only=True, read_only=True)
    except Exception as e:
        raise ParseError(f"Excel 文件无法打开（{type(e).__name__}）") from e

    segments: list[Segment] = []
    try:
        for index, sheet in enumerate(workbook.worksheets, start=1):
            markdown = _sheet_to_markdown(sheet)
            if markdown:
                segments.append(
                    Segment(f"工作表：{sheet.title}\n{markdown}", index, TABLE)
                )
    finally:
        workbook.close()
    if not segments:
        raise ParseError("Excel 中没有可提取的数据")
    return segments


def _sheet_to_markdown(sheet) -> str:
    rows: list[list[str]] = []
    for row in sheet.iter_rows(values_only=True):
        cells = [_cell_text(value) for value in row]
        if not any(cells):
            continue
        rows.append(cells)
        if len(rows) >= settings.kb_sheet_max_rows:
            logger.info("工作表 %s 超过 %d 行，截断", sheet.title, len(rows))
            break
    if not rows:
        return ""
    width = max(len(row) for row in rows)
    rows = [row + [""] * (width - len(row)) for row in rows]
    # 空表头补列名：Markdown 表格没有列名就无法读懂每列是什么
    rows[0] = [cell or f"列{i + 1}" for i, cell in enumerate(rows[0])]
    return _rows_to_markdown(rows)


def _cell_text(value) -> str:
    if value is None:
        return ""
    if isinstance(value, float) and value.is_integer():
        return str(int(value))
    return _escape_cell(str(value))
